"""Reading a client's risk and control matrix, whatever its columns are called.

A client RCM is untrusted input. It is parsed to plain cell text and nothing else:
no formula, macro, link, or embedded object is evaluated, and nothing is passed to a
model. Columns are matched to the platform's fields by a reviewed synonym list, and a
person confirms the match. Every column the platform has no field for is kept, under
the client's own heading. The client's terms for area, frequency, and nature are kept
too, next to the standard value a person maps each one to.
"""

import csv
import io
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from ai_workforce_api.controls import DOMAINS, FREQUENCIES, NATURES

MAX_BYTES = 5_000_000
MAX_UNZIPPED = 60_000_000
MAX_SHEETS = 20
MAX_ROWS = 2_000
MAX_COLUMNS = 80
MAX_CELL = 4_000
MAX_CONTROLS = 1_000
MAX_DISTINCT = 200
HEADER_SCAN = 25

# The fields the platform understands. Everything else is a custom column.
FIELDS: tuple[tuple[str, str], ...] = (
    ("control_ref", "Control ID"),
    ("title", "Control name"),
    ("activity", "Control description"),
    ("risk", "Risk"),
    ("domain", "Process or area"),
    ("frequency", "Frequency"),
    ("nature", "Nature"),
    ("key_control", "Key control"),
    ("owner_name", "Control owner"),
    ("owner_email", "Owner email"),
    ("applications", "Applications"),
    ("test_attributes", "Test procedures"),
    ("pbc_items", "Evidence"),
)
FIELD_KEYS = tuple(key for key, _label in FIELDS)
CUSTOM = "custom"
IGNORE = "ignore"
CHOICES = (*FIELD_KEYS, CUSTOM, IGNORE)
REQUIRED = ("control_ref", "activity")
# Fields whose client values are mapped to a standard value.
VALUE_FIELDS: dict[str, tuple[str, ...]] = {
    "domain": DOMAINS,
    "frequency": FREQUENCIES,
    "nature": NATURES,
}

# Reviewed header synonyms, compared after normalizing case and punctuation.
SYNONYMS: dict[str, tuple[str, ...]] = {
    "control_ref": (
        "control id",
        "control ref",
        "control reference",
        "control number",
        "control no",
        "control code",
        "control",
        "ctrl",
        "ctrl id",
        "ctrl ref",
        "ctrl no",
        "itgc id",
        "itgc ref",
        "ref",
        "reference",
        "id",
        "key control id",
    ),
    "title": (
        "control name",
        "control title",
        "title",
        "name",
        "short description",
        "control summary",
        "control short name",
        "control objective",
        "objective",
    ),
    "activity": (
        "control description",
        "control activity",
        "control activities",
        "description",
        "activity",
        "control narrative",
        "control procedure",
        "control statement",
        "control detail",
        "control details",
    ),
    "risk": (
        "risk",
        "risks",
        "risk description",
        "risk statement",
        "what could go wrong",
        "wcgw",
        "risk of material misstatement",
        "rmm",
        "risk addressed",
    ),
    "domain": (
        "domain",
        "itgc domain",
        "control domain",
        "process",
        "process area",
        "sub process",
        "subprocess",
        "area",
        "control area",
        "category",
        "itgc category",
        "control category",
        "cycle",
        "business process",
        "business cycle",
        "process name",
        "domain process",
    ),
    "frequency": (
        "frequency",
        "control frequency",
        "freq",
        "how often",
        "timing",
        "periodicity",
    ),
    "nature": (
        "nature",
        "control nature",
        "control type",
        "type",
        "manual automated",
        "manual or automated",
        "automated manual",
        "automation",
        "method",
        "control method",
    ),
    "key_control": (
        "key",
        "key control",
        "key control y n",
        "key non key",
        "key or non key",
        "is key",
        "key yes no",
    ),
    "owner_name": (
        "owner",
        "control owner",
        "owner name",
        "control owner name",
        "process owner",
        "responsible",
        "responsible party",
        "responsible person",
        "performer",
        "control performer",
        "control operator",
    ),
    "owner_email": (
        "owner email",
        "control owner email",
        "email",
        "e mail",
        "contact email",
        "owner e mail",
    ),
    "applications": (
        "application",
        "applications",
        "system",
        "systems",
        "application system",
        "in scope application",
        "in scope applications",
        "in scope systems",
        "it system",
        "it systems",
        "supporting system",
        "supporting systems",
        "technology",
        "platform",
    ),
    "test_attributes": (
        "test procedures",
        "test procedure",
        "testing procedures",
        "testing procedure",
        "test steps",
        "testing steps",
        "test attributes",
        "attributes",
        "audit procedures",
        "testing approach",
        "test approach",
    ),
    "pbc_items": (
        "evidence",
        "evidence requested",
        "required evidence",
        "evidence required",
        "pbc",
        "pbc items",
        "pbc list",
        "documentation",
        "supporting documentation",
        "key reports",
        "ipe",
    ),
}

# Reviewed keyword rules for the client's terms. The first matching rule wins.
VALUE_RULES: dict[str, tuple[tuple[tuple[str, ...], str], ...]] = {
    "frequency": (
        (
            ("multiple times", "several times", "continuous", "real time", "hourly", "recurring"),
            "RECURRING",
        ),
        (("semi annual", "semiannual", "bi annual", "biannual", "half year"), "QUARTERLY"),
        (("daily", "day"), "DAILY"),
        (("week",), "WEEKLY"),
        (("month",), "MONTHLY"),
        (("quarter", "qtr"), "QUARTERLY"),
        (("annual", "year"), "ANNUAL"),
        (
            (
                "as needed",
                "as required",
                "ad hoc",
                "adhoc",
                "event",
                "occurrence",
                "upon",
                "per change",
                "per request",
                "per release",
                "release cycle",
                "on demand",
                "when",
                "transaction",
            ),
            "EVENT_DRIVEN",
        ),
    ),
    "nature": (
        (
            ("dependent", "itdm", "hybrid", "semi", "partial", "assisted", "it manual"),
            "IT_DEPENDENT_MANUAL",
        ),
        (("automat", "system", "configur", "application control"), "AUTOMATED"),
        (("manual",), "MANUAL"),
    ),
    # Business words come before IT words, so "Exchange rates" is treasury, not change.
    "domain": (
        (
            (
                "entity level",
                "entity wide",
                "company level",
                "elc",
                "governance",
                "tone at the top",
                "coso",
                "whistleblow",
                "code of conduct",
                "anti fraud",
            ),
            "ENTITY",
        ),
        (
            ("application control", "itac", "automated application", "it application"),
            "ITAC",
        ),
        (
            (
                "ipe",
                "key report",
                "information produced",
                "report logic",
                "end user computing",
                "euc",
            ),
            "IPE",
        ),
        (
            (
                "order to cash",
                "o2c",
                "otc",
                "revenue",
                "sales",
                "billing",
                "receivable",
                "credit management",
                "cash receipt",
                "cash application",
                "customer",
            ),
            "ORDER_TO_CASH",
        ),
        (
            (
                "procure",
                "p2p",
                "ptp",
                "purchas",
                "payable",
                "vendor",
                "supplier",
                "expenditure",
                "disbursement",
            ),
            "PROCURE_TO_PAY",
        ),
        (
            (
                "hire to retire",
                "h2r",
                "htr",
                "payroll",
                "human resource",
                "hr",
                "compensation",
                "benefits",
            ),
            "HIRE_TO_RETIRE",
        ),
        (
            (
                "treasury",
                "cash management",
                "cash",
                "bank",
                "investment",
                "debt",
                "exchange",
                "foreign currency",
                "fx",
                "liquidity",
            ),
            "TREASURY",
        ),
        (
            (
                "inventory",
                "cost of sales",
                "cost of goods",
                "cogs",
                "warehouse",
                "production",
                "manufactur",
                "costing",
                "stock",
            ),
            "INVENTORY",
        ),
        (
            (
                "fixed asset",
                "ppe",
                "pp e",
                "property plant",
                "lease",
                "capex",
                "capital expenditure",
                "depreciation",
            ),
            "FIXED_ASSETS",
        ),
        (("tax", "vat", "gst"), "TAX"),
        (
            (
                "record to report",
                "r2r",
                "rtr",
                "financial close",
                "period end",
                "close",
                "journal",
                "general ledger",
                "gl",
                "consolidation",
                "financial reporting",
                "financial statement",
            ),
            "RECORD_TO_REPORT",
        ),
        (
            (
                "access",
                "security",
                "user",
                "password",
                "logical",
                "identity",
                "provisioning",
                "authentication",
                "privileged",
            ),
            "ACCESS",
        ),
        (
            (
                "change",
                "development",
                "sdlc",
                "program",
                "release",
                "implementation",
                "patch",
                "acquisition",
            ),
            "CHANGE",
        ),
        (
            (
                "operation",
                "backup",
                "job",
                "batch",
                "incident",
                "schedul",
                "recovery",
                "monitoring",
                "interface",
                "computer",
                "physical",
                "data center",
                "datacenter",
                "vendor management",
                "third party",
                "service organization",
                "soc 1",
                "soc report",
            ),
            "OPERATIONS",
        ),
    ),
}

# What a category column says about a control: which areas its process can be in. A
# category that is itself an area (ITAC, key reports, entity-level) decides outright.
IT_AREAS = ("ACCESS", "CHANGE", "OPERATIONS")
BUSINESS_AREAS = (
    "ENTITY",
    "ORDER_TO_CASH",
    "PROCURE_TO_PAY",
    "HIRE_TO_RETIRE",
    "TREASURY",
    "INVENTORY",
    "FIXED_ASSETS",
    "TAX",
    "RECORD_TO_REPORT",
    "OTHER_BUSINESS",
)
_CATEGORY_RULES: tuple[tuple[tuple[str, ...], str], ...] = (
    (("itac", "application control", "it application"), "ITAC"),
    (("ipe", "key report", "information produced"), "IPE"),
    (("elc", "entity level", "entity wide", "company level"), "ENTITY"),
    (("itgc", "general control", "it general", "general computer"), "IT"),
    (("business process", "business cycle", "process control", "bp", "financial"), "BUSINESS"),
)
# Words that make a column the control's type (preventive or detective), never its nature.
_CONTROL_TYPE_WORDS = ("preventive", "preventative", "detective", "corrective", "directive")

DEFAULT_ATTRIBUTE = "The control operated as described in the control activity."
DEFAULT_EVIDENCE = "Evidence that the control operated for each selected sample."

_EMAIL = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_BULLET = re.compile(r"^\s*(?:[-*•▪●]|\(?[0-9]{1,2}[.)]|\(?[a-zA-Z][.)])\s+")
# Credential-shaped values only. Control text talks about passwords and keys all the
# time ("password length = 8", "risk-based"), so a broader pattern would drop real rows.
_CREDENTIAL = re.compile(
    r"(?i)(?:api[_-]?key|client[_-]?secret|refresh[_-]?token|access[_-]?token|private[_-]?key)"
    r"\s*[:=]\s*\S{8,}"
    r"|password\s*[:=]\s*(?=\S*[0-9])(?=\S*[A-Za-z])\S{6,}"
    r"|(?<![A-Za-z0-9])sk-[A-Za-z0-9_-]{20,}"
    r"|BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY"
    r"|eyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}"
)
_YES = {"yes", "y", "true", "key", "x", "1", "key control", "k"}


class RcmFileError(Exception):
    def __init__(self, detail: str) -> None:
        self.detail = detail


@dataclass(frozen=True)
class Sheet:
    name: str
    rows: list[list[str]]


@dataclass
class ParsedControl:
    row: int
    control_ref: str
    title: str
    activity: str
    risk: str
    domain: str
    domain_label: str
    frequency: str
    frequency_label: str
    nature: str
    nature_label: str
    key_control: bool
    owner_name: str
    owner_email: str
    applications: list[str]
    test_attributes: list[str]
    pbc_items: list[str]
    custom_fields: list[tuple[str, str]]


@dataclass
class Skipped:
    row: int
    reason: str


@dataclass
class ValueChoice:
    value: str
    count: int
    suggested: str | None


@dataclass
class BuildResult:
    controls: list[ParsedControl] = field(default_factory=list)
    skipped: list[Skipped] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


# Reading files


def parse_file(file_name: str, content: bytes) -> tuple[str, list[Sheet]]:
    """Read an upload to plain cell text. Returns the file kind and its sheets."""
    if not content:
        raise RcmFileError("file is empty")
    if len(content) > MAX_BYTES:
        raise RcmFileError("file is too large")
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        raise RcmFileError("save the workbook as .xlsx or .csv")
    if content.startswith(b"PK\x03\x04"):
        return "XLSX", _xlsx(content)
    if file_name.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raise RcmFileError("file is not a readable workbook")
    return "CSV", [_csv(file_name, content)]


def _xlsx(content: bytes) -> list[Sheet]:
    # Imported here so a CSV-only path never loads the workbook reader.
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5_000 or sum(item.file_size for item in entries) > MAX_UNZIPPED:
                raise RcmFileError("workbook is too large")
        # Values only: cached results instead of formulas, and no external links.
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except RcmFileError:
        raise
    except Exception as error:
        raise RcmFileError("file is not a readable workbook") from error
    sheets: list[Sheet] = []
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            rows: list[list[str]] = []
            for values in worksheet.iter_rows(values_only=True):
                if len(rows) >= MAX_ROWS + HEADER_SCAN:
                    break
                rows.append([_cell(value) for value in values[:MAX_COLUMNS]])
            sheets.append(
                Sheet(name=_clean(str(worksheet.title))[:100] or "Sheet", rows=_trim(rows))
            )
    except Exception as error:
        raise RcmFileError("file is not a readable workbook") from error
    finally:
        workbook.close()
    return _usable(sheets)


def _csv(file_name: str, content: bytes) -> Sheet:
    if b"\x00" in content[:4096]:
        raise RcmFileError("file is not a CSV or workbook")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    sample = text[:8192]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    rows: list[list[str]] = []
    try:
        for values in csv.reader(io.StringIO(text), delimiter=delimiter):
            if len(rows) >= MAX_ROWS + HEADER_SCAN:
                break
            rows.append([_cell(value) for value in values[:MAX_COLUMNS]])
    except csv.Error as error:
        raise RcmFileError("file is not a readable CSV") from error
    name = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", file_name)[:100] or "Sheet"
    return _usable([Sheet(name=_clean(name) or "Sheet", rows=_trim(rows))])[0]


def _usable(sheets: list[Sheet]) -> list[Sheet]:
    if not any(sum(1 for row in sheet.rows if any(row)) >= 2 for sheet in sheets):
        raise RcmFileError("file has no table of controls")
    return sheets


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean(str(value))[:MAX_CELL]


def _clean(text: str) -> str:
    text = _CONTROL_CHARS.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
    return "\n".join(" ".join(line.split()) for line in text.split("\n")).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(rows[-1]):
        rows.pop()
    width = max(
        (max((i + 1 for i, v in enumerate(row) if v), default=0) for row in rows), default=0
    )
    return [(row + [""] * width)[:width] for row in rows]


# Layout


def normalize(header: str) -> str:
    text = header.lower().replace("(s)", "s")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


# Single words distinctive enough to name a field inside a longer heading.
_PHRASE_WORDS = {"risk", "frequency", "evidence", "nature", "domain", "applications", "attributes"}


def _synonym_field(header: str) -> tuple[str, int] | None:
    """The field a header names, and how strongly: 2 for an exact synonym, 1 for a phrase."""
    text = normalize(header)
    if text == "":
        return None
    for key in FIELD_KEYS:
        if text in SYNONYMS[key]:
            return key, 2
    best: tuple[str, int, int] | None = None
    padded = f" {text} "
    for key in FIELD_KEYS:
        for synonym in SYNONYMS[key]:
            if (" " in synonym or synonym in _PHRASE_WORDS) and f" {synonym} " in padded:
                if best is None or len(synonym) > best[2]:
                    best = (key, 1, len(synonym))
    return None if best is None else (best[0], best[1])


def detect_header(rows: list[list[str]]) -> int:
    """The 1-based row that holds the column headings."""
    best_row, best_score = 1, -1
    for index, row in enumerate(rows[:HEADER_SCAN]):
        cells = [cell for cell in row if cell]
        if len(cells) < 2:
            continue
        short = sum(1 for cell in cells if len(cell) <= 80 and not _numeric(cell))
        named = sum(3 for cell in cells if _synonym_field(cell) is not None)
        score = short + named
        if score > best_score:
            best_row, best_score = index + 1, score
    return best_row


def best_sheet(sheets: list[Sheet]) -> int:
    def score(sheet: Sheet) -> tuple[int, int]:
        header = detect_header(sheet.rows)
        headings = sheet.rows[header - 1] if sheet.rows else []
        named = sum(1 for cell in headings if _synonym_field(cell) is not None)
        return named, len(sheet.rows)

    return max(range(len(sheets)), key=lambda index: score(sheets[index]))


def headers(rows: list[list[str]], header_row: int) -> list[str]:
    raw = rows[header_row - 1] if 0 < header_row <= len(rows) else []
    width = max((len(row) for row in rows), default=0)
    names: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        cell = raw[index] if index < len(raw) else ""
        name = " ".join(cell.split())[:120] or f"Column {_letters(index)}"
        count = seen.get(name.lower(), 0) + 1
        seen[name.lower()] = count
        names.append(name if count == 1 else f"{name} ({count})")
    return names


def data_rows(rows: list[list[str]], header_row: int) -> list[tuple[int, list[str]]]:
    """Rows below the header, with their 1-based row numbers, blank rows left out."""
    return [
        (header_row + offset + 1, row)
        for offset, row in enumerate(rows[header_row : header_row + MAX_ROWS])
        if any(row)
    ]


# Fields whose values are distinctive enough to name a column without a heading.
_VALUE_NAMED = ("nature", "frequency", "key_control")

# The area may come from two columns: a category (ITGC, Business Process, Key Report)
# and the process itself (Order to cash, Program change). They are read together.
DOMAIN_COLUMNS = 2


def column_values(rows: list[list[str]], header_row: int, limit: int = 60) -> list[list[str]]:
    """The distinct non-blank values in each column, in the order they appear."""
    width = max((len(row) for row in rows), default=0)
    values: list[list[str]] = [[] for _ in range(width)]
    seen: list[set[str]] = [set() for _ in range(width)]
    for _number, row in data_rows(rows, header_row):
        for index, cell in enumerate(row[:width]):
            text = " ".join(cell.split())
            key = text.lower()
            if text and key not in seen[index] and len(values[index]) < limit:
                seen[index].add(key)
                values[index].append(text)
    return values


def _id_like(value: str) -> bool:
    text = value.strip()
    return len(text) <= 30 and any(ch.isdigit() for ch in text) and len(text.split()) <= 3


def looks_like_control_type(values: list[str]) -> bool:
    """Preventive or detective: the control's type, which is not its nature."""
    if not values:
        return False
    hits = sum(1 for v in values if any(w in v.lower() for w in _CONTROL_TYPE_WORDS))
    return hits / len(values) >= 0.6


def _share(values: list[str], fits: Callable[[str], bool]) -> int:
    """2 when most values fit, 1 when a good share do, else 0."""
    if not values:
        return 0
    share = sum(1 for v in values if fits(v)) / len(values)
    return 2 if share >= 0.7 else 1 if share >= 0.35 else 0


def value_fit(key: str, values: list[str]) -> int:
    """How well a column's values look like the field: 0 (no), 1 (somewhat), 2 (clearly)."""
    if not values:
        return 0
    if key == "nature":
        if looks_like_control_type(values):
            return 0
        return _share(values, lambda v: suggest_value("nature", v) is not None)
    if key == "frequency":
        return _share(values, lambda v: suggest_value("frequency", v) is not None)
    if key == "domain":
        # A category such as ITGC or Business Process is an area column too.
        return _share(
            values,
            lambda v: (
                suggest_value("domain", v) is not None
                or _rule(normalize(v), _CATEGORY_RULES) is not None
            ),
        )
    if key == "key_control":
        return _share(values, lambda v: normalize(v) in _YES | {"no", "n", "false", "non key", "0"})
    if key == "control_ref":
        unique = len({v.lower() for v in values}) == len(values)
        return _share(values, _id_like) if unique else 0
    if key in ("risk", "activity"):
        # Prose: a sentence, not an identifier or a one-word code.
        if _share(values, _id_like) == 2:
            return 0
        return _share(values, lambda v: len(v) >= 25 and " " in v)
    if key == "title":
        if _share(values, _id_like) == 2:
            return 0
        return _share(values, lambda v: 8 <= len(v) < 120)
    return 0


def suggest_fields(
    names: list[str],
    remembered: dict[str, str] | None = None,
    samples: list[list[str]] | None = None,
) -> list[str]:
    """One choice per column, from the headings and the values together.

    A heading names a field; the values confirm or veto it: a "Nature" column holding
    preventive and detective is the control's type, not its nature, and "Risk ID" is an
    identifier, not the risk. Each field goes to one column, except the area, which can
    come from a category column and a process column together.
    """
    memory = remembered or {}
    values = samples if samples is not None else [[] for _ in names]
    choices = [CUSTOM] * len(names)
    counts: dict[str, int] = {}
    for index, name in enumerate(names):
        known = memory.get(normalize(name))
        if known in FIELD_KEYS and counts.get(known, 0) < (
            DOMAIN_COLUMNS if known == "domain" else 1
        ):
            choices[index] = known
            counts[known] = counts.get(known, 0) + 1
        elif known == IGNORE:
            choices[index] = IGNORE
    candidates: list[tuple[int, int, int, int, str]] = []
    for index, name in enumerate(names):
        if choices[index] != CUSTOM or normalize(name) in memory:
            continue
        column = values[index] if index < len(values) else []
        if looks_like_control_type(column):
            continue
        match = _synonym_field(name)
        for key in FIELD_KEYS:
            strength = match[1] if match is not None and match[0] == key else 0
            fit = value_fit(key, column)
            if column and fit == 0 and key in (*VALUE_FIELDS, "key_control", "risk"):
                # The values say otherwise. Another column that fits wins over a phrase
                # in this heading, and one that fits clearly wins over an exact heading.
                better = max(
                    (value_fit(key, other) for j, other in enumerate(values) if j != index),
                    default=0,
                )
                if (strength < 2 and better >= 1) or (strength == 2 and better == 2):
                    continue
            score = 2 * strength + 2 * fit
            # Values alone name a column only where the vocabulary is distinctive.
            if score >= 2 and (strength > 0 or (fit == 2 and key in _VALUE_NAMED)):
                candidates.append((score, strength, -index, index, key))
    for _score, _strength, _order, index, key in sorted(candidates, reverse=True):
        if choices[index] != CUSTOM:
            continue
        if counts.get(key, 0) >= (DOMAIN_COLUMNS if key == "domain" else 1):
            continue
        choices[index] = key
        counts[key] = counts.get(key, 0) + 1
    return choices


def field_columns(column_fields: list[str], key: str) -> list[int]:
    """The columns mapped to a field: one, or two for the area."""
    return [index for index, choice in enumerate(column_fields) if choice == key]


def field_value(row: list[str], column_fields: list[str], key: str) -> str:
    """A row's value for a field. The area's two columns read as "category / process"."""
    parts = [
        " ".join(row[index].split())
        for index in field_columns(column_fields, key)
        if index < len(row) and row[index].strip()
    ]
    return " / ".join(parts)


def check_fields(column_fields: list[str], width: int) -> list[str]:
    if len(column_fields) != width or any(choice not in CHOICES for choice in column_fields):
        raise ValueError("column mapping is invalid")
    mapped = [choice for choice in column_fields if choice in FIELD_KEYS]
    if any(mapped.count(key) > (DOMAIN_COLUMNS if key == "domain" else 1) for key in set(mapped)):
        raise ValueError("a field is mapped twice")
    labels = dict(FIELDS)
    return [f"Choose the column for {labels[key]}." for key in REQUIRED if key not in mapped]


# Values


def _rule(text: str, rules: tuple[tuple[tuple[str, ...], str], ...]) -> str | None:
    padded = f" {text} "
    for words, result in rules:
        if any(word in padded if len(word) > 3 else f" {word} " in padded for word in words):
            return result
    return None


def suggest_value(field_key: str, value: str) -> str | None:
    text = normalize(value)
    if text == "":
        return None
    canonical = text.upper().replace(" ", "_")
    if canonical in VALUE_FIELDS[field_key]:
        return canonical
    if field_key == "domain":
        return _suggest_domain(value)
    return _rule(text, VALUE_RULES[field_key])


def _suggest_domain(value: str) -> str | None:
    """The area for a term, or for "category / process" read from two columns.

    A category that is an area (ITAC, key report, entity-level) decides. One that is a
    kind (ITGC, business process) narrows the process to that kind's areas, so "ITGC /
    Vendor management" is IT operations, not procure to pay.
    """
    parts = [normalize(part) for part in value.split(" / ") if normalize(part)]
    if not parts:
        return None
    kind: str | None = None
    process: list[str] = []
    for part in parts:
        category = _rule(part, _CATEGORY_RULES)
        if category in ("ITAC", "IPE", "ENTITY"):
            return category
        if category is not None and _rule(part, VALUE_RULES["domain"]) is None:
            kind = category
        else:
            if category is not None:
                kind = category
            process.append(part)
    allowed = IT_AREAS if kind == "IT" else BUSINESS_AREAS if kind == "BUSINESS" else None
    rules = (
        VALUE_RULES["domain"]
        if allowed is None
        else tuple(rule for rule in VALUE_RULES["domain"] if rule[1] in allowed)
    )
    for part in process or parts:
        found = _rule(part, rules)
        if found is not None:
            return found
    return _rule(" ".join(parts), rules)


def value_key(value: str) -> str:
    return " ".join(value.split()).lower()


def value_choices(
    rows: list[list[str]], header_row: int, column_fields: list[str]
) -> dict[str, list[ValueChoice]]:
    """The distinct client values for each mapped field, with a suggested standard value."""
    result: dict[str, list[ValueChoice]] = {}
    body = data_rows(rows, header_row)
    # A row without a control ID is skipped at import, so its terms need no mapping.
    if "control_ref" in column_fields:
        ref = column_fields.index("control_ref")
        body = [(number, row) for number, row in body if ref < len(row) and row[ref].strip()]
    for key in VALUE_FIELDS:
        counts: dict[str, ValueChoice] = {}
        for _number, row in body:
            label = field_value(row, column_fields, key)[:120]
            item = counts.get(value_key(label))
            if item is None:
                if len(counts) >= MAX_DISTINCT:
                    raise ValueError(f"{key} has too many distinct values")
                counts[value_key(label)] = ValueChoice(label, 1, suggest_value(key, label))
            else:
                item.count += 1
        result[key] = sorted(counts.values(), key=lambda item: (-item.count, item.value.lower()))
    return result


def check_values(value_maps: dict[str, dict[str, str]]) -> None:
    if set(value_maps) - set(VALUE_FIELDS):
        raise ValueError("value mapping is invalid")
    for key, mapping in value_maps.items():
        if len(mapping) > MAX_DISTINCT or any(
            len(raw) > 120 or chosen not in VALUE_FIELDS[key] for raw, chosen in mapping.items()
        ):
            raise ValueError("value mapping is invalid")


# Building controls


def build(
    rows: list[list[str]],
    header_row: int,
    column_fields: list[str],
    value_maps: dict[str, dict[str, str]],
) -> BuildResult:
    """Turn mapped rows into controls. Rows that cannot be used are listed, not guessed."""
    names = headers(rows, header_row)
    result = BuildResult()
    result.problems.extend(check_fields(column_fields, len(names)))
    choices = value_choices(rows, header_row, column_fields)
    labels = dict(FIELDS)
    for key, items in choices.items():
        unmapped = [
            item for item in items if value_maps.get(key, {}).get(value_key(item.value)) is None
        ]
        if unmapped:
            shown = ", ".join(f'"{item.value or "(blank)"}"' for item in unmapped[:3])
            more = f" and {len(unmapped) - 3} more" if len(unmapped) > 3 else ""
            result.problems.append(f"Map {labels[key].lower()} values {shown}{more}.")
    if result.problems:
        return result
    column = {choice: index for index, choice in enumerate(column_fields) if choice in FIELD_KEYS}
    seen: dict[str, int] = {}
    bad_emails = 0
    for number, row in data_rows(rows, header_row):

        def cell(key: str, row: list[str] = row) -> str:
            index = column.get(key)
            return row[index] if index is not None and index < len(row) else ""

        ref = " ".join(cell("control_ref").split())
        activity = cell("activity").strip()
        title = " ".join(cell("title").split())
        if ref == "":
            result.skipped.append(Skipped(number, "No control ID."))
            continue
        if len(ref) > 60:
            result.skipped.append(Skipped(number, "Control ID is longer than 60 characters."))
            continue
        if ref.lower() in seen:
            result.skipped.append(Skipped(number, f"Control ID repeats row {seen[ref.lower()]}."))
            continue
        if activity == "":
            activity = title
        if activity == "":
            result.skipped.append(Skipped(number, "No control description."))
            continue
        if any(_CREDENTIAL.search(value) for value in row):
            result.skipped.append(Skipped(number, "The row looks like it holds a secret."))
            continue
        if len(result.controls) >= MAX_CONTROLS:
            result.skipped.append(Skipped(number, f"Only {MAX_CONTROLS} controls are imported."))
            continue
        seen[ref.lower()] = number
        email = cell("owner_email").strip().lower()
        if email != "" and _EMAIL.fullmatch(email) is None:
            bad_emails += 1
            email = ""
        owner = " ".join(cell("owner_name").split())[:200]
        if email == "" and _EMAIL.fullmatch(owner.lower()):
            email, owner = owner.lower(), ""
        mapped: dict[str, tuple[str, str]] = {}
        for key in VALUE_FIELDS:
            label = field_value(row, column_fields, key)[:120]
            mapped[key] = (value_maps[key][value_key(label)], label)
        key_text = normalize(cell("key_control"))
        result.controls.append(
            ParsedControl(
                row=number,
                control_ref=ref,
                title=(title or _summary(activity))[:200],
                activity=activity[:4000],
                risk=cell("risk").strip()[:4000],
                domain=mapped["domain"][0],
                domain_label=mapped["domain"][1],
                frequency=mapped["frequency"][0],
                frequency_label=mapped["frequency"][1],
                nature=mapped["nature"][0],
                nature_label=mapped["nature"][1],
                key_control=True if "key_control" not in column else key_text in _YES,
                owner_name=owner,
                owner_email=email[:200],
                applications=_split(cell("applications"), r"[,;\n|]", 20, 120),
                test_attributes=_split(cell("test_attributes"), r"\n|;", 20, 500)
                or [DEFAULT_ATTRIBUTE],
                pbc_items=_split(cell("pbc_items"), r"\n|;", 20, 500) or [DEFAULT_EVIDENCE],
                custom_fields=[
                    (names[index], row[index][:MAX_CELL])
                    for index, choice in enumerate(column_fields)
                    if choice == CUSTOM and index < len(row) and row[index] != ""
                ],
            )
        )
    if bad_emails:
        result.warnings.append(
            f"{bad_emails} owner email{'s are' if bad_emails > 1 else ' is'} not valid and "
            "will be left blank."
        )
    if not result.controls:
        result.problems.append("No row has a control ID and a description.")
    return result


def _split(text: str, pattern: str, limit: int, width: int) -> list[str]:
    items: list[str] = []
    for part in re.split(pattern, text):
        item = " ".join(_BULLET.sub("", part).split())[:width]
        if item and item not in items:
            items.append(item)
    return items[:limit]


def _summary(activity: str) -> str:
    first = re.split(r"(?<=[.!?])\s", " ".join(activity.split()), maxsplit=1)[0]
    return first if len(first) <= 120 else first[:117].rstrip() + "..."


def _numeric(text: str) -> bool:
    return re.fullmatch(r"[-+]?[0-9.,%/: ]+", text) is not None


def _letters(index: int) -> str:
    name = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(ord("A") + remainder) + name
    return name
