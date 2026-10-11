"""Reading stored evidence to plain text. Deterministic, values only, and bounded.

Nothing here calls a model, a network, or another program. Spreadsheets give their
cell values (cached results, never formulas or links), documents their text, and PDFs
their text layer. A screenshot, a scanned page, or an image pasted into a document
becomes a normalized PNG for OCR, which the analysis job hands to the firm's engine.

The file is untrusted. Every reader has a size limit, archives are checked before they
are opened, and secret-looking values are redacted from the text that is kept.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from typing import Any, Literal

from ai_workforce_api.inbound import clean_text

MAX_TEXT = 200_000
MAX_PDF_PAGES = 300
MAX_OCR_IMAGES = 10
MAX_SCAN_ROWS = 100_000
MAX_COLUMNS = 60
MAX_SHEETS = 10
MAX_ENTRIES = 5_000
MAX_UNZIPPED = 150_000_000
MAX_PIXELS = 40_000_000
# A page with less text than this is treated as a scan.
SCANNED_PAGE_CHARS = 25
OCR_LONG_EDGE = 3_000
OCR_SHORT_EDGE = 1_000
MIN_IMAGE_EDGE = 40

Method = Literal["TEXT", "TABLE", "DOCUMENT", "PDF_TEXT", "OCR", "NONE"]
Status = Literal["EXTRACTED", "PARTIAL", "NEEDS_OCR", "UNSUPPORTED", "UNREADABLE"]
FileKind = Literal["pdf", "workbook", "csv", "document", "text", "image", "other"]

_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_IMAGE_SUFFIX = re.compile(r"(?i)\.(png|jpe?g|gif|bmp|tiff?|webp)$")
_TEXT_SUFFIX = re.compile(r"(?i)\.(txt|log|text|md)$")
_CSV_SUFFIX = re.compile(r"(?i)\.(csv|tsv)$")
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


@dataclass(frozen=True)
class Extraction:
    file_kind: FileKind
    method: Method
    status: Status
    text: str
    truncated: bool
    redacted: bool
    pages: int | None
    rows: int | None
    columns: tuple[str, ...]
    # PNG images still to be read by OCR, in page order.
    images: tuple[bytes, ...]
    note: str


def extract(content: bytes, file_name: str, content_type: str) -> Extraction:
    """Read one file. Never raises for a bad file: the result says what happened."""
    kind = sniff(content, file_name, content_type)
    try:
        if kind == "pdf":
            return _pdf(content)
        if kind == "workbook":
            return _workbook(content)
        if kind == "document":
            return _document(content)
        if kind == "csv":
            return _csv(content)
        if kind == "text":
            return _text(content)
        if kind == "image":
            return _image(content)
    except _Unreadable as error:
        return _none(kind, "UNREADABLE", error.note)
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        return _none(
            "other",
            "UNSUPPORTED",
            "Older Office files are not read. Ask for .xlsx, .docx, or PDF.",
        )
    return _none("other", "UNSUPPORTED", "This file type is not read. A person reviews it.")


def sniff(content: bytes, file_name: str, content_type: str) -> FileKind:
    """The file's kind from its bytes first. The name and declared type only break ties."""
    head = content[:16]
    if head.startswith(b"%PDF-"):
        return "pdf"
    if (
        head.startswith(b"\x89PNG\r\n\x1a\n")
        or head.startswith(b"\xff\xd8\xff")
        or head.startswith((b"GIF87a", b"GIF89a"))
        or head.startswith(b"BM")
        or head.startswith((b"II*\x00", b"MM\x00*"))
        or (head.startswith(b"RIFF") and content[8:12] == b"WEBP")
    ):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist()[:MAX_ENTRIES])
        except zipfile.BadZipFile:
            return "other"
        if "xl/workbook.xml" in names:
            return "workbook"
        if "word/document.xml" in names:
            return "document"
        return "other"
    if head.startswith(b"\xd0\xcf\x11\xe0") or b"\x00" in content[:4096]:
        return "other"
    if _CSV_SUFFIX.search(file_name) or content_type.lower() in {"text/csv", "application/csv"}:
        return "csv"
    if _TEXT_SUFFIX.search(file_name) or content_type.lower().startswith("text/plain"):
        return "text"
    if _IMAGE_SUFFIX.search(file_name):
        return "other"
    return "text" if _decodes(content) else "other"


def with_ocr(extraction: Extraction, texts: list[str], note: str = "") -> Extraction:
    """Add OCR text to an extraction that was waiting for it."""
    read = [text.strip() for text in texts]
    found = [text for text in read if text]
    notes = " ".join(part for part in (extraction.note, note) if part)
    if extraction.file_kind == "image":
        combined = found[0] if found else ""
    elif extraction.file_kind == "pdf":
        pages = [f"[Page {index + 1}, read by OCR]\n{text}" for index, text in enumerate(read)]
        combined = "\n\n".join(part for part in (extraction.text, *pages) if part.strip())
    else:
        images = [f"[Image {index + 1}, read by OCR]\n{text}" for index, text in enumerate(read)]
        combined = "\n\n".join(part for part in (extraction.text, *images) if part.strip())
    if not found and not extraction.text.strip():
        return replace(
            extraction,
            status="UNREADABLE",
            images=(),
            method="OCR",
            note=_join(notes, "OCR found no text. A person reviews the file."),
        )
    cleaned = clean_text(combined, MAX_TEXT)
    method: Method = "OCR" if extraction.file_kind in {"image", "pdf"} else extraction.method
    # A note here is a reading limit, such as pages left unread.
    complete = len(found) == len(read) and not extraction.note
    status: Status = "EXTRACTED" if complete else "PARTIAL"
    return replace(
        extraction,
        method=method,
        status=status,
        text=cleaned.text,
        truncated=cleaned.truncated or extraction.truncated,
        redacted=cleaned.redacted or extraction.redacted,
        images=(),
        note=notes[:300],
    )


def without_ocr(extraction: Extraction, note: str) -> Extraction:
    """No OCR engine could read the images. Keep any text, and say a person must read it."""
    if extraction.text.strip():
        return replace(extraction, status="PARTIAL", images=(), note=_join(extraction.note, note))
    return replace(extraction, status="NEEDS_OCR", images=(), note=_join(extraction.note, note))


# Readers


class _Unreadable(Exception):
    def __init__(self, note: str) -> None:
        self.note = note


def _pdf(content: bytes) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            # Many PDFs are encrypted with an empty password only to restrict printing.
            if not reader.decrypt(""):
                raise _Unreadable("The PDF is password-protected. Ask for an open copy.")
        pages = list(reader.pages)
    except _Unreadable:
        raise
    except (PyPdfError, ValueError, KeyError, TypeError, OSError) as error:
        raise _Unreadable("The PDF could not be read.") from error
    count = len(pages)
    texts: list[str] = []
    for page in pages[:MAX_PDF_PAGES]:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
    scanned = [index for index, text in enumerate(texts) if len(text.strip()) < SCANNED_PAGE_CHARS]
    note = f"Only the first {MAX_PDF_PAGES} pages were read." if count > MAX_PDF_PAGES else ""
    readable = "\n\n".join(
        f"[Page {index + 1}]\n{text.strip()}"
        for index, text in enumerate(texts)
        if len(text.strip()) >= SCANNED_PAGE_CHARS
    )
    images: list[bytes] = []
    if scanned:
        for index in scanned[:MAX_OCR_IMAGES]:
            image = _page_image(pages[index])
            if image is not None:
                images.append(image)
        if len(scanned) > MAX_OCR_IMAGES:
            note = _join(note, f"Only the first {MAX_OCR_IMAGES} scanned pages are read by OCR.")
    cleaned = clean_text(readable, MAX_TEXT)
    if images:
        status: Status = "NEEDS_OCR"
    elif scanned and not readable:
        return _none("pdf", "UNREADABLE", "The PDF has no text and no page images.", pages=count)
    else:
        status = "PARTIAL" if scanned or note else "EXTRACTED"
    return Extraction(
        file_kind="pdf",
        method="PDF_TEXT",
        status=status,
        text=cleaned.text,
        truncated=cleaned.truncated,
        redacted=cleaned.redacted,
        pages=count,
        rows=None,
        columns=(),
        images=tuple(images),
        note=note,
    )


def _page_image(page: Any) -> bytes | None:
    """The largest image on a scanned page, ready for OCR."""
    best: Any = None
    best_area = 0
    try:
        for item in list(page.images)[:20]:
            picture = item.image
            if picture is None:
                continue
            area = int(picture.width) * int(picture.height)
            if area > best_area:
                best, best_area = picture, area
    except Exception:
        return None
    if best is None:
        return None
    try:
        return _ocr_png(best)
    except _Unreadable:
        return None


def _workbook(content: bytes) -> Extraction:
    sheets, note = _workbook_sheets(content)
    return _table("workbook", sheets, note)


def _workbook_sheets(content: bytes) -> tuple[list[tuple[str, list[list[str]]]], str]:
    from openpyxl import load_workbook

    _check_archive(content)
    try:
        # Values only: cached results instead of formulas, and no external links.
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except Exception as error:
        raise _Unreadable("The workbook could not be read.") from error
    sheets: list[tuple[str, list[list[str]]]] = []
    note = ""
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            note = f"Only the first {MAX_SHEETS} sheets were read."
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            rows: list[list[str]] = []
            for values in worksheet.iter_rows(values_only=True):
                if len(rows) >= MAX_SCAN_ROWS:
                    note = _join(note, f"Only the first {MAX_SCAN_ROWS:,} rows were read.")
                    break
                rows.append([_cell(value) for value in values[:MAX_COLUMNS]])
            sheets.append((_line(str(worksheet.title))[:100] or "Sheet", _trim(rows)))
    except Exception as error:
        raise _Unreadable("The workbook could not be read.") from error
    finally:
        workbook.close()
    return sheets, note


def _csv(content: bytes) -> Extraction:
    sheets, note = _csv_sheets(content)
    return _table("csv", sheets, note)


def _csv_sheets(content: bytes) -> tuple[list[tuple[str, list[list[str]]]], str]:
    text = _decode(content)
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    rows: list[list[str]] = []
    note = ""
    try:
        for values in csv.reader(io.StringIO(text), delimiter=delimiter):
            if len(rows) >= MAX_SCAN_ROWS:
                note = f"Only the first {MAX_SCAN_ROWS:,} rows were read."
                break
            rows.append([_line(value)[:2000] for value in values[:MAX_COLUMNS]])
    except csv.Error as error:
        raise _Unreadable("The CSV could not be read.") from error
    return [("Sheet", _trim(rows))], note


@dataclass(frozen=True)
class Table:
    """The main sheet of a population: its header and every data row, as cell text."""

    sheet: str
    columns: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    note: str


def read_table(content: bytes, file_name: str, content_type: str) -> Table | None:
    """The largest sheet of a workbook or CSV, for sampling. None for any other file."""
    kind = sniff(content, file_name, content_type)
    try:
        if kind == "workbook":
            sheets, note = _workbook_sheets(content)
        elif kind == "csv":
            sheets, note = _csv_sheets(content)
        else:
            return None
    except _Unreadable:
        return None
    filled = [(name, rows) for name, rows in sheets if any(any(row) for row in rows)]
    if not filled:
        return None
    name, rows = max(filled, key=lambda sheet: sum(1 for row in sheet[1] if any(row)))
    header_index = next(index for index, row in enumerate(rows) if any(row))
    header = rows[header_index]
    width = len(header)
    data = tuple(
        tuple((row + [""] * width)[:width]) for row in rows[header_index + 1 :] if any(row)
    )
    return Table(sheet=name, columns=tuple(cell[:200] for cell in header), rows=data, note=note)


def _table(kind: FileKind, sheets: list[tuple[str, list[list[str]]]], note: str) -> Extraction:
    """Sheets as text. The largest sheet gives the columns and the row count."""
    filled = [(name, rows) for name, rows in sheets if any(any(row) for row in rows)]
    if not filled:
        return _none(kind, "UNREADABLE", "The file has no values.")
    main = max(filled, key=lambda sheet: sum(1 for row in sheet[1] if any(row)))[1]
    header_index = next(index for index, row in enumerate(main) if any(row))
    columns = tuple(cell[:80] for cell in main[header_index] if cell)[:MAX_COLUMNS]
    rows = sum(1 for row in main[header_index + 1 :] if any(row))
    parts: list[str] = []
    size = 0
    for name, sheet_rows in filled:
        lines = [f"## Sheet: {name}"] if len(filled) > 1 or kind == "workbook" else []
        lines += [" | ".join(row) for row in sheet_rows if any(row)]
        block = "\n".join(lines)
        parts.append(block)
        size += len(block)
        if size > MAX_TEXT:
            break
    cleaned = clean_text("\n\n".join(parts), MAX_TEXT)
    return Extraction(
        file_kind=kind,
        method="TABLE",
        status="PARTIAL" if note else "EXTRACTED",
        text=cleaned.text,
        truncated=cleaned.truncated,
        redacted=cleaned.redacted,
        pages=None,
        rows=rows,
        columns=columns,
        images=(),
        note=note,
    )


def _document(content: bytes) -> Extraction:
    from defusedxml import ElementTree

    _check_archive(content)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            body = archive.read("word/document.xml")
            media = [
                item
                for item in archive.infolist()
                if item.filename.startswith("word/media/") and _IMAGE_SUFFIX.search(item.filename)
            ]
            pictures = [archive.read(item) for item in media[: MAX_OCR_IMAGES * 2]]
        root = ElementTree.fromstring(body)
    except Exception as error:
        raise _Unreadable("The document could not be read.") from error
    lines: list[str] = []
    container = root.find(f"{_WORD}body")
    for block in list(container) if container is not None else []:
        if block.tag == f"{_WORD}p":
            lines.append(_paragraph(block))
        elif block.tag == f"{_WORD}tbl":
            for row in block.iter(f"{_WORD}tr"):
                cells = [
                    " ".join(_paragraph(p) for p in cell.iter(f"{_WORD}p")).strip()
                    for cell in row.iter(f"{_WORD}tc")
                ]
                lines.append(" | ".join(cells))
    cleaned = clean_text("\n".join(lines), MAX_TEXT)
    images: list[bytes] = []
    for picture in pictures:
        if len(images) >= MAX_OCR_IMAGES:
            break
        try:
            images.append(_ocr_png_bytes(picture))
        except _Unreadable:
            continue
    note = (
        f"Only the first {MAX_OCR_IMAGES} images are read by OCR."
        if len(media) > MAX_OCR_IMAGES
        else ""
    )
    if not cleaned.text and not images:
        return _none("document", "UNREADABLE", "The document has no text.")
    return Extraction(
        file_kind="document",
        method="DOCUMENT",
        status="NEEDS_OCR" if images else "EXTRACTED",
        text=cleaned.text,
        truncated=cleaned.truncated,
        redacted=cleaned.redacted,
        pages=None,
        rows=None,
        columns=(),
        images=tuple(images),
        note=note,
    )


def _paragraph(element: Any) -> str:
    parts: list[str] = []
    for node in element.iter():
        if node.tag == f"{_WORD}t" and node.text:
            parts.append(node.text)
        elif node.tag == f"{_WORD}tab":
            parts.append("\t")
        elif node.tag in (f"{_WORD}br", f"{_WORD}cr"):
            parts.append("\n")
    return "".join(parts)


def _text(content: bytes) -> Extraction:
    cleaned = clean_text(_decode(content), MAX_TEXT)
    if not cleaned.text:
        return _none("text", "UNREADABLE", "The file is empty.")
    return Extraction(
        file_kind="text",
        method="TEXT",
        status="EXTRACTED",
        text=cleaned.text,
        truncated=cleaned.truncated,
        redacted=cleaned.redacted,
        pages=None,
        rows=None,
        columns=(),
        images=(),
        note="",
    )


def _image(content: bytes) -> Extraction:
    return Extraction(
        file_kind="image",
        method="NONE",
        status="NEEDS_OCR",
        text="",
        truncated=False,
        redacted=False,
        pages=1,
        rows=None,
        columns=(),
        images=(_ocr_png_bytes(content),),
        note="",
    )


# Images


def _ocr_png_bytes(content: bytes) -> bytes:
    from PIL import Image, UnidentifiedImageError

    try:
        picture = Image.open(io.BytesIO(content))
        width, height = picture.size
    except (UnidentifiedImageError, OSError, ValueError) as error:
        raise _Unreadable("The image could not be read.") from error
    if width * height > MAX_PIXELS:
        raise _Unreadable("The image is too large to read.")
    try:
        picture.seek(0)
        picture.load()
    except (OSError, ValueError, Image.DecompressionBombError) as error:
        raise _Unreadable("The image could not be read.") from error
    return _ocr_png(picture)


def _ocr_png(picture: Any) -> bytes:
    """Grayscale PNG, upright, scaled so text is a readable size for OCR."""
    from PIL import Image, ImageOps

    try:
        width, height = picture.size
        if width * height > MAX_PIXELS:
            raise _Unreadable("The image is too large to read.")
        if min(width, height) < MIN_IMAGE_EDGE:
            raise _Unreadable("The image is too small to read.")
        upright = ImageOps.exif_transpose(picture) or picture
        gray = upright.convert("L")
        long_edge = max(gray.size)
        if long_edge > OCR_LONG_EDGE:
            gray.thumbnail((OCR_LONG_EDGE, OCR_LONG_EDGE), Image.Resampling.LANCZOS)
        elif long_edge < OCR_SHORT_EDGE:
            scale = min(3.0, OCR_SHORT_EDGE / long_edge)
            gray = gray.resize(
                (int(gray.width * scale), int(gray.height * scale)), Image.Resampling.LANCZOS
            )
        output = io.BytesIO()
        gray.save(output, format="PNG")
    except _Unreadable:
        raise
    except Exception as error:
        raise _Unreadable("The image could not be read.") from error
    return output.getvalue()


# Helpers


def _check_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as error:
        raise _Unreadable("The file is not a readable archive.") from error
    if len(entries) > MAX_ENTRIES or sum(item.file_size for item in entries) > MAX_UNZIPPED:
        raise _Unreadable("The file expands to more than can be read safely.")


def _none(kind: FileKind, status: Status, note: str, *, pages: int | None = None) -> Extraction:
    return Extraction(
        file_kind=kind,
        method="NONE",
        status=status,
        text="",
        truncated=False,
        redacted=False,
        pages=pages,
        rows=None,
        columns=(),
        images=(),
        note=note,
    )


def _decodes(content: bytes) -> bool:
    try:
        content[:8192].decode("utf-8")
    except UnicodeDecodeError as error:
        # A sample can end inside a character.
        return error.start >= 8188
    return True


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _line(str(value))[:2000]


def _line(text: str) -> str:
    return " ".join(_CONTROL_CHARS.sub(" ", text).split())


def _trim(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def _join(first: str, second: str) -> str:
    return " ".join(part for part in (first, second) if part)[:300]
