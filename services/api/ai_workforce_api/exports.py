"""Engagement exports (Milestone 8.20): records to Excel or CSV, from the same data the
agent answers with (engagement_data).

A file is built from one read of the records, checked (its rows are the records the
filters chose, every column is present), stored, and only then recorded and offered for
download, so a link never points at a file that is not there. Text from clients and
owners goes into the file as text: a cell is never a formula. A file can be downloaded for
RETENTION_DAYS by people who can see its engagement; after that its bytes are removed and
the record of the export stays. Making and downloading a file are audited; neither changes
the engagement's records.
"""

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any, Literal
from uuid import UUID, uuid4

import psycopg
from fastapi import FastAPI, HTTPException, Request, Response
from psycopg.types.json import Json
from pydantic import BaseModel

from ai_workforce_api import engagement_data as data
from ai_workforce_api.authorization import (
    Action,
    NotVisible,
    Principal,
    _call,
    _session_principal,
    authorize,
)
from ai_workforce_api.db import transaction
from ai_workforce_api.settings import Settings
from ai_workforce_api.storage import DeletableStore, ObjectStore, StoreError, digest_of, open_store

RETENTION_DAYS = 30
Dataset = Literal["evidence_requests", "controls", "exceptions", "review_items"]
FileFormat = Literal["xlsx", "csv"]
StoreFactory = Callable[[Settings, psycopg.Connection[Any]], ObjectStore]

CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
# Records go to a spreadsheet; reports (8.21) to a document.
RECORD_FORMATS = ("xlsx", "csv")
REPORT_FORMATS = ("pdf", "docx")
TITLES: dict[str, str] = {
    "evidence_requests": "Evidence requests",
    "controls": "Control testing status",
    "exceptions": "Exceptions",
    "review_items": "Open review items",
}
# Each dataset's columns: the heading, and the record field it shows.
COLUMNS: dict[str, list[tuple[str, str]]] = {
    "evidence_requests": [
        ("Reference", "reference"),
        ("Control", "control_ref"),
        ("Control title", "control_title"),
        ("Application", "application"),
        ("Process", "process"),
        ("Testing phase", "phase"),
        ("Evidence type", "evidence_type"),
        ("Request", "description"),
        ("Control owner", "control_owner"),
        ("Recipient", "recipient"),
        ("Requested on", "requested_on"),
        ("Due on", "due_on"),
        ("Status", "status"),
        ("Status group", "status_group"),
        ("Overdue", "overdue"),
        ("Days overdue", "days_overdue"),
        ("Files received", "files_received"),
        ("Evidence review", "evidence_review"),
        ("Follow-up", "follow_up"),
        ("Last updated (UTC)", "last_updated"),
    ],
    "controls": [
        ("Control", "control_ref"),
        ("Title", "title"),
        ("Application", "application"),
        ("Process", "process"),
        ("Owner", "owner"),
        ("Key control", "key_control"),
        ("Frequency", "frequency"),
        ("Testing status", "testing_status"),
        ("Stage", "stage"),
        ("Test sheet", "test_sheet"),
        ("Test result", "test_result"),
        ("Conclusion", "conclusion"),
        ("Failed", "failed"),
        ("Independent review", "review"),
        ("Open review notes", "open_review_notes"),
        ("Workpapers", "workpapers"),
        ("Evidence requests", "evidence_requests"),
        ("Evidence pending", "evidence_pending"),
        ("Evidence overdue", "evidence_overdue"),
        ("Open deficiencies", "open_deficiencies"),
        ("Next step", "next_step"),
        ("Waiting on", "waiting_on"),
    ],
    "exceptions": [
        ("Exception", "exception_id"),
        ("Kind", "kind"),
        ("Control", "control_ref"),
        ("Application", "application"),
        ("Title", "title"),
        ("Description", "description"),
        ("Type", "type"),
        ("Severity (as recorded)", "severity_as_recorded"),
        ("Status", "status"),
        ("Remediation owner", "remediation_owner"),
        ("Remediation due", "remediation_due"),
        ("Identified on", "identified_on"),
        ("Decision pending", "pending_decision"),
    ],
    "review_items": [
        ("Item", "item"),
        ("Kind", "kind"),
        ("Control", "control_ref"),
        ("Application", "application"),
        ("About", "about"),
        ("Detail", "detail"),
        ("Severity (as recorded)", "severity_as_recorded"),
        ("Status", "status"),
        ("Raised by", "raised_by"),
        ("Since", "since"),
        ("Days open", "days_open"),
        ("Response", "response"),
    ],
}
ROWS_KEY = {
    "evidence_requests": "requests",
    "controls": "controls",
    "exceptions": "exceptions",
    "review_items": "items",
}


class ExportFailed(Exception):
    """The file could not be made or stored. Nothing was recorded or offered."""


class ExportBody(BaseModel):
    id: UUID
    dataset: str
    file_format: str
    file_name: str
    rows: int
    filters: dict[str, Any]
    as_of: datetime
    expires_at: datetime


@dataclass(frozen=True)
class Extract:
    """The records chosen for a file, and everything said about how they were chosen."""

    dataset: str
    result: dict[str, Any]
    rows: list[dict[str, Any]]


def extract(
    connection: psycopg.Connection[Any],
    principal: Principal,
    organization_id: UUID,
    project_id: UUID,
    dataset: str,
    filters: dict[str, Any],
) -> Extract:
    """Every record the filters choose (no page), from one read, checked to add up."""
    if dataset == "evidence_requests":
        result = data.evidence_requests(
            connection,
            principal,
            organization_id,
            project_id,
            data.RequestFilters(**filters),
            sort="most_overdue",
            limit=None,
        )
    elif dataset == "controls":
        result = data.controls(
            connection,
            principal,
            organization_id,
            project_id,
            data.ControlFilters(**filters),
            limit=None,
        )
    elif dataset == "exceptions":
        result = data.exceptions(
            connection, principal, organization_id, project_id, limit=None, **filters
        )
    elif dataset == "review_items":
        result = data.review_items(
            connection, principal, organization_id, project_id, limit=None, **filters
        )
    else:
        raise ExportFailed(f"There is no export of {dataset}.")
    rows = list(result[ROWS_KEY[dataset]])
    # The file holds exactly the records counted, each with every column.
    if len(rows) != result["total"]:
        raise ExportFailed("The records changed while they were read. Try again.")
    for row in rows:
        missing = [key for _heading, key in COLUMNS[dataset] if key not in row]
        if missing:
            raise ExportFailed(f"A record is missing {', '.join(missing)}.")
    return Extract(dataset, result, rows)


# Rendering

_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def _cell(value: Any) -> Any:
    """A value as a cell: text stays text (never a formula), flags read Yes and No."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, dict):
        value = "; ".join(f"{key}: {count}" for key, count in value.items())
    text = str(value)
    return f"'{text}" if text.startswith(_FORMULA) else text


def _about(extract: Extract, generated_by: str, generated_at: datetime) -> list[tuple[str, str]]:
    result = extract.result
    filters = result.get("filters") or {}
    lines = [
        ("Engagement", str(result["engagement"])),
        ("Client", str(result["client"])),
        ("Report", TITLES[extract.dataset]),
        (
            "Filters",
            "; ".join(f"{key.replace('_', ' ')} = {value}" for key, value in filters.items())
            or "None: every record",
        ),
        ("Records", f"{len(extract.rows):,}"),
        ("Data as of (UTC)", str(result["as_of"])),
        ("Generated (UTC)", generated_at.isoformat(timespec="seconds")),
        ("Generated by", generated_by),
        ("Rules", str(result["rules"])),
        (
            "Source",
            "The engagement's records in AI Workforce, read at one moment. Counts, statuses, "
            "and overdue days are worked out by the application, not written by AI. "
            '"Not recorded" means the record does not hold that value; it is not zero.',
        ),
    ]
    for check in result.get("data_quality", []):
        lines.append(
            (
                "Data check",
                f"{check['finding']}: {check['count']} ({', '.join(check['examples'])})",
            )
        )
    return lines


def render_xlsx(extract: Extract, generated_by: str, generated_at: datetime) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    about = book.active
    if about is None:
        raise ExportFailed("The workbook could not be made.")
    about.title = "About"
    about.append([TITLES[extract.dataset]])
    about["A1"].font = Font(bold=True, size=14)
    about.append([])
    for label, value in _about(extract, generated_by, generated_at):
        about.append([label, _cell(value)])
    about.column_dimensions["A"].width = 20
    about.column_dimensions["B"].width = 100
    for row in about.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        if len(row) > 1:
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet = book.create_sheet("Records")
    columns = COLUMNS[extract.dataset]
    sheet.append([heading for heading, _key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for record in extract.rows:
        sheet.append([_cell(record.get(key)) for _heading, key in columns])
    sheet.freeze_panes = "A2"
    if extract.rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(extract.rows) + 1}"
    for index, (heading, key) in enumerate(columns, start=1):
        longest = max([len(heading), *(len(str(_cell(r.get(key)))) for r in extract.rows[:500])])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 60)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def render_csv(extract: Extract) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    columns = COLUMNS[extract.dataset]
    writer.writerow([heading for heading, _key in columns])
    for record in extract.rows:
        writer.writerow([_cell(record.get(key)) for _heading, key in columns])
    # With a byte order mark, so Excel opens the accents and dashes as written.
    return ("﻿" + out.getvalue()).encode("utf-8")


def _file_name(engagement: str, dataset: str, as_of: datetime, file_format: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "-", engagement).strip("-")[:60] or "engagement"
    return f"{stem}-{dataset.replace('_', '-')}-{as_of:%Y%m%d-%H%M}Z.{file_format}"


# Making and fetching


def create_export(
    connection: psycopg.Connection[Any],
    settings: Settings,
    principal: Principal,
    organization_id: UUID,
    project_id: UUID,
    dataset: str,
    file_format: str,
    filters: dict[str, Any],
    store_for: StoreFactory = open_store,
    narrative: str = "",
) -> ExportBody:
    """Make the file, store it, and record it, in that order: no record, no link."""
    from ai_workforce_api import reports

    report = dataset in reports.REPORTS
    allowed = REPORT_FORMATS if report else RECORD_FORMATS
    if file_format not in allowed:
        kinds = "PDF (pdf) or Word (docx)" if report else "Excel (xlsx) or CSV"
        raise ExportFailed(f"This is made as {kinds}, not {file_format}.")
    who = connection.execute(
        "SELECT coalesce(nullif(display_name, ''), email::text) FROM users WHERE id = %s",
        (principal.id,),
    ).fetchone()
    by = str(who[0]) if who else ""
    now = datetime.now(UTC)
    try:
        if report:
            built = reports.build(
                connection, principal, organization_id, project_id, dataset, filters, by,
                narrative,
            )  # fmt: skip
            engagement, rows = built.engagement, built.rows
            as_of = datetime.fromisoformat(built.as_of)
            applied = {"scope": built.scope, **filters}
            content = (
                reports.render_pdf(built) if file_format == "pdf" else reports.render_docx(built)
            )
        else:
            chosen = extract(connection, principal, organization_id, project_id, dataset, filters)
            engagement, rows = str(chosen.result["engagement"]), len(chosen.rows)
            as_of = datetime.fromisoformat(str(chosen.result["as_of"]))
            applied = dict(chosen.result.get("filters") or {})
            content = render_xlsx(chosen, by, now) if file_format == "xlsx" else render_csv(chosen)
    except ExportFailed, NotVisible:
        raise
    except reports.ReportFailed as error:
        raise ExportFailed(str(error)) from None
    except data.DataUnavailable as error:
        raise ExportFailed(str(error)) from None
    except Exception as error:
        raise ExportFailed("The file could not be made. Nothing was saved.") from error
    export_id = uuid4()
    name = _file_name(engagement, dataset, as_of, file_format)
    key = f"exports/{export_id}"
    store = store_for(settings, connection)
    try:
        digest = store.put(key, content)
    except StoreError as error:
        raise ExportFailed("The file could not be stored. Nothing was saved; try again.") from error
    if digest != digest_of(content):
        raise ExportFailed("The stored file does not match. Nothing was saved; try again.")
    with transaction(connection) as cursor:
        row = cursor.execute(
            """
            WITH stamp AS (SELECT clock_timestamp() AS now)
            INSERT INTO engagement_exports (
                id, organization_id, project_id, dataset, file_format, filters, row_count,
                as_of, file_name, object_key, content_sha256, byte_size, created_by,
                narrative, created_at, expires_at
            )
            -- One timestamp for both, so the retention is exactly the stated days (two
            -- clock_timestamp() calls could differ by a microsecond).
            SELECT %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                   stamp.now, stamp.now + make_interval(days => %s)
            FROM stamp
            RETURNING expires_at
            """,
            (
                export_id, organization_id, project_id, dataset, file_format, Json(applied),
                rows, as_of, name, key, digest, len(content), principal.id,
                " ".join(narrative.split())[:4000], RETENTION_DAYS,
            ),
        ).fetchone()  # fmt: skip
        _audit(
            cursor, organization_id, project_id, principal.id, "export.created", export_id,
            {"dataset": dataset, "format": file_format, "rows": rows,
             "filters": applied, "as_of": as_of.isoformat(),
             "narrative": bool(narrative.strip())},
        )  # fmt: skip
    purge_expired(connection, store, project_id)
    if row is None:
        raise ExportFailed("The file could not be recorded.")
    return ExportBody(
        id=export_id,
        dataset=dataset,
        file_format=file_format,
        file_name=name,
        rows=rows,
        filters=applied,
        as_of=as_of,
        expires_at=row[0],
    )


def purge_expired(connection: psycopg.Connection[Any], store: ObjectStore, project_id: UUID) -> int:
    """Remove the bytes of the engagement's expired exports; their records stay."""
    expired = connection.execute(
        """
        SELECT id, object_key FROM engagement_exports
        WHERE project_id = %s AND purged_at IS NULL AND expires_at < clock_timestamp()
        """,
        (project_id,),
    ).fetchall()
    removed = 0
    for export_id, key in expired:
        if isinstance(store, DeletableStore):
            try:
                store.delete(str(key))
            except StoreError:
                continue
        with transaction(connection) as cursor:
            cursor.execute(
                "UPDATE engagement_exports SET purged_at = clock_timestamp() WHERE id = %s",
                (export_id,),
            )
        removed += 1
    return removed


class ExportGone(Exception):
    """The export is past its retention; its file is no longer kept."""


def fetch_export(
    connection: psycopg.Connection[Any],
    settings: Settings,
    principal: Principal,
    organization_id: UUID,
    project_id: UUID,
    export_id: UUID,
    store_for: StoreFactory = open_store,
) -> tuple[str, str, bytes]:
    """The file, for someone who can see its engagement, while it is kept. Audited."""
    authorize(
        connection,
        principal=principal,
        action=Action.TASK_READ,
        organization_id=organization_id,
        project_id=project_id,
    )
    row = connection.execute(
        """
        SELECT file_name, file_format, object_key, content_sha256,
               expires_at < clock_timestamp() OR purged_at IS NOT NULL
        FROM engagement_exports
        WHERE id = %s AND project_id = %s AND organization_id = %s
        """,
        (export_id, project_id, organization_id),
    ).fetchone()
    if row is None:
        raise NotVisible()
    if row[4]:
        raise ExportGone()
    content = store_for(settings, connection).get(str(row[2]))
    if digest_of(content) != row[3]:
        raise ExportFailed("The stored file does not match its record.")
    with transaction(connection) as cursor:
        _audit(
            cursor, organization_id, project_id, principal.id, "export.downloaded", export_id,
            {},
        )  # fmt: skip
    return str(row[0]), CONTENT_TYPES[str(row[1])], content


def _audit(
    cursor: psycopg.Cursor[Any],
    organization_id: UUID,
    project_id: UUID,
    actor_id: UUID,
    action: str,
    export_id: UUID,
    payload: dict[str, Any],
) -> None:
    cursor.execute(
        """
        INSERT INTO audit_events (organization_id, project_id, actor_user_id, action,
                                  target_type, target_id, payload)
        VALUES (%s, %s, %s, %s, 'export', %s, %s)
        """,
        (organization_id, project_id, actor_id, action, export_id, Json(payload)),
    )


def register_exports(
    app: FastAPI, settings: Settings, store_for: StoreFactory = open_store
) -> None:
    @app.get("/organizations/{organization_id}/projects/{project_id}/exports/{export_id}")
    def download(
        organization_id: UUID, project_id: UUID, export_id: UUID, request: Request
    ) -> Response:
        principal = _session_principal(request)
        try:
            name, content_type, content = _call(
                settings,
                lambda connection: fetch_export(
                    connection, settings, principal, organization_id, project_id, export_id,
                    store_for,
                ),
            )  # fmt: skip
        except ExportGone:
            raise HTTPException(
                status_code=410, detail="this export has expired; make it again"
            ) from None
        except ExportFailed, StoreError:
            raise HTTPException(status_code=503, detail="export is unavailable") from None
        safe = re.sub(r"[^A-Za-z0-9._-]+", "-", name)
        return Response(
            content=content,
            media_type=content_type,
            headers={
                "Content-Disposition": f'attachment; filename="{safe}"',
                "Cache-Control": "no-store",
            },
        )
