"""Engagement reports (Milestone 8.21): PDF and Word, from the engagement's records.

Five reports share one path. A builder reads the records through engagement_data (the
same figures the agent answers with) and lays them out as a Report: sections of facts,
lists, and tables. The report is checked (every table holds exactly the records counted,
and it says what it covers and when the records were read), then drawn as PDF or Word.

What the records say and what the application worked out fill the report. The only
AI-written part is an optional narrative, printed in its own section and marked as an
interpretation that is not verified. Severity is shown only as recorded, and the report
never states an audit opinion or a conclusion the records do not hold.
"""

import io
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any, Literal
from uuid import UUID
from xml.sax.saxutils import escape

import psycopg

from ai_workforce_api import engagement_data as data
from ai_workforce_api.authorization import Principal

ReportKind = Literal[
    "status_report",
    "evidence_tracker",
    "control_summary",
    "exceptions_report",
    "executive_summary",
]
REPORTS: dict[str, str] = {
    "status_report": "Engagement status report",
    "evidence_tracker": "Pending evidence tracker",
    "control_summary": "Control testing summary",
    "exceptions_report": "Exceptions and remediation report",
    "executive_summary": "Executive summary",
}
PENDING = ["draft", "awaiting_approval", "with_client", "partly_received"]
NARRATIVE_LABEL = "AI-written summary"
NARRATIVE_NOTE = (
    "Written by the AI agent from the figures in this report. It is an interpretation, "
    "not verified, and not an audit conclusion."
)
STATUS_WORDS = {
    "complete": "Complete",
    "awaiting_review": "Awaiting review",
    "in_testing": "In testing",
    "awaiting_evidence": "Awaiting evidence",
    "not_started": "Not started",
}


class ReportFailed(Exception):
    """The report could not be made, or did not add up. Nothing was saved."""


@dataclass
class Table:
    title: str
    headers: list[str]
    rows: list[list[str]]
    # How many records the data layer counted: the table must hold exactly these.
    total: int | None = None
    note: str = ""
    # Relative column widths; equal when left out.
    widths: list[float] | None = None
    empty: str = "None."
    # Print the title above the table, where the section holds more than the table.
    titled: bool = False


@dataclass
class Section:
    title: str
    intro: str = ""
    facts: list[tuple[str, str]] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    kind: Literal["records", "narrative", "about"] = "records"


@dataclass
class Report:
    kind: str
    title: str
    engagement: str
    client: str
    as_of: str
    rules: str
    scope: str
    generated_by: str
    generated_at: datetime
    sections: list[Section]
    rows: int
    landscape: bool = False


def _when(value: str | None) -> str:
    return "Not recorded" if not value else str(value)


def _date(value: Any) -> str:
    return "Not recorded" if value in (None, "") else str(value)


def _clip(text: str, limit: int) -> str:
    """Long text cut at a word, marked as cut."""
    if len(text) <= limit:
        return text
    head = text[:limit]
    if not text[limit].isspace():
        head = head.rsplit(" ", 1)[0]
    return head.rstrip(" ,;:") + "…"


def _words(value: str) -> str:
    """A recorded code as words: NOT_CONCLUDED -> Not concluded."""
    return value.replace("_", " ").capitalize()


def _moment(value: str) -> str:
    """The data-as-of time as people read it: 2026-10-16 21:03:30 UTC."""
    try:
        moment = datetime.fromisoformat(value)
    except ValueError:
        return f"{value} UTC"
    return f"{moment.astimezone(UTC):%Y-%m-%d %H:%M:%S} UTC"


def _scope_line(filters: dict[str, Any]) -> str:
    parts = []
    for key, value in filters.items():
        if value in (None, "", False, []):
            continue
        name = key.replace("_", " ")
        if value is True:
            parts.append(name)
        elif isinstance(value, list):
            parts.append(f"{name}: {' or '.join(str(item).replace('_', ' ') for item in value)}")
        else:
            parts.append(f"{name}: {value}")
    return "; ".join(parts) or "Every record"


def _about(result: dict[str, Any], scope: str, narrative: bool) -> Section:
    lines = [
        "Source: the engagement's records in AI Workforce, read at one moment (data as of "
        f"{_moment(str(result['as_of']))}). Counts, statuses, days overdue, and completion are "
        "worked out by the application, not written by AI.",
        f"Scope: {scope}. Cancelled requests and out-of-scope controls are left out.",
        f"Rules: {result['rules']} (docs/metrics.md).",
        '"Not recorded" means the record does not hold the value; it is not zero.',
        "Severity is shown only as the team recorded it. This report states no audit "
        "opinion: the only conclusions in it are those recorded and signed off in the "
        "engagement.",
    ]
    if narrative:
        lines.append(f'The section "{NARRATIVE_LABEL}" is the only AI-written part.')
    return Section("About this report", bullets=lines, kind="about")


def _checks(checks: list[dict[str, Any]]) -> Section | None:
    if not checks:
        return None
    return Section(
        "Data checks",
        intro="Records that contradict each other or miss what the work needs. They are "
        "reported as they are, not corrected.",
        bullets=[f"{c['finding']}: {c['count']} ({', '.join(c['examples'])})" for c in checks],
    )


def _request_table(
    rows: list[dict[str, Any]],
    total: int,
    title: str,
    empty: str = "No evidence requests match.",
    titled: bool = False,
) -> Table:
    return Table(
        title,
        [
            "Request", "Control", "Application", "Phase", "Evidence type", "Owner",
            "Requested", "Due", "Status", "Days overdue", "Follow-up",
        ],
        [
            [
                f"{r['reference']}: {_clip(r['description'], 140)}",
                f"{r['control_ref']} {r['control_title']}",
                r["application"],
                r["phase"].title() if r["phase"] != "Not recorded" else r["phase"],
                r["evidence_type"],
                r["control_owner"],
                r["requested_on"],
                r["due_on"],
                data.GROUP_LABELS.get(r["status_group"], r["status"]),
                str(r["days_overdue"]) if r["overdue"] else "",
                r["follow_up"],
            ]
            for r in rows
        ],
        total=total,
        widths=[2.8, 2.0, 1.2, 0.9, 1.5, 1.2, 1.25, 1.25, 1.1, 1.1, 1.5],
        empty=empty,
        titled=titled,
    )  # fmt: skip


def _control_table(rows: list[dict[str, Any]], total: int) -> Table:
    return Table(
        "Controls",
        [
            "Control", "Application", "Process", "Testing status", "Evidence pending",
            "Test result", "Workpapers", "Review", "Open notes", "Deficiencies", "Next step",
        ],
        [
            [
                f"{r['control_ref']} {r['title']}",
                r["application"],
                r["process"],
                STATUS_WORDS.get(r["testing_status"], r["testing_status"]),
                f"{r['evidence_pending']} of {r['evidence_requests']}"
                + (
                    " (" + ", ".join(
                        f"{phase.title() if phase != 'Not recorded' else phase} {count}"
                        for phase, count in r["evidence_pending_by_phase"].items()
                    ) + ")"
                    if r["evidence_pending_by_phase"]
                    else ""
                ),
                _words(r["test_result"]),
                r["workpapers"],
                _words(r["review"]),
                str(r["open_review_notes"]),
                str(r["open_deficiencies"]),
                r["next_step"] or "",
            ]
            for r in rows
        ],
        total=total,
        widths=[2.8, 1.2, 1.3, 1.2, 1.5, 1.2, 1.3, 1.0, 0.9, 1.2, 1.9],
        empty="No controls match.",
    )  # fmt: skip


def _exception_action(row: dict[str, Any]) -> str:
    if row["pending_decision"]:
        return "A decision is pending"
    return {
        "OPEN": "Agree the remediation",
        "REMEDIATION_IN_PROGRESS": "Remediate",
        "READY_FOR_RETEST": "Retest",
        "NOT_EVALUATED": "Auditor to judge whether it is a deficiency",
    }.get(row["status"], "")


def _exception_table(rows: list[dict[str, Any]], total: int) -> Table:
    return Table(
        "Exceptions",
        [
            "Exception", "Control", "Application", "Description", "Severity (as recorded)",
            "Status", "Remediation owner", "Target date", "Outstanding",
        ],
        [
            [
                f"{r['exception_id']}: {r['title']}",
                r["control_ref"],
                r["application"],
                _clip(r["description"], 300),
                _words(r["severity_as_recorded"])
                if r["severity_as_recorded"] != "Not recorded"
                else r["severity_as_recorded"],
                _words(r["status"]),
                r["remediation_owner"],
                _date(r["remediation_due"]),
                _exception_action(r),
            ]
            for r in rows
        ],
        total=total,
        widths=[2.2, 1.2, 1.2, 3.4, 1.3, 1.2, 1.3, 1.0, 1.6],
        empty="No open exceptions.",
    )  # fmt: skip


def _review_table(rows: list[dict[str, Any]], total: int) -> Table:
    return Table(
        "Outstanding review items",
        ["Item", "Control", "About", "Status", "Since", "Days open"],
        [
            [
                r["item"],
                r["control_ref"],
                _clip(r["detail"], 200),
                r["status"],
                _date(r["since"]),
                "" if r["days_open"] is None else str(r["days_open"]),
            ]
            for r in rows
        ],
        total=total,
        widths=[1.4, 1.4, 4.0, 1.3, 1.0, 0.8],
        empty="No open review notes, and no prepared workpaper waits for review.",
    )


def _overview_facts(status: dict[str, Any]) -> list[tuple[str, str]]:
    controls, evidence = status["controls"], status["evidence"]
    completion = controls["completion"]
    percent = "n/a" if completion["percent"] is None else f"{completion['percent']}%"
    return [
        ("Engagement status", str(status["engagement_status"]).title()),
        ("Period end", _date(status["period_end"])),
        ("Controls in scope", f"{controls['in_scope']} ({controls['key']} key)"),
        ("Testing complete", f"{completion['complete']} of {completion['in_scope']} ({percent})"),
        ("Failed testing", str(controls["failed"])),
        ("Evidence requests pending", str(evidence["pending"])),
        ("Overdue", str(evidence["overdue"])),
        ("Received, not yet reviewed", str(evidence["received_not_reviewed"])),
        ("Open deficiencies", str(status["exceptions"]["open_deficiencies"])),
        ("Open review notes", str(status["review"]["open_review_notes"])),
        ("Missed milestones", str(status["schedule"]["missed_milestones"])),
    ]


def _next_actions(status: dict[str, Any], requests: dict[str, Any]) -> list[str]:
    """What the records say needs doing next, with the records behind each item."""
    actions = []
    late = [r for r in requests["requests"] if r["overdue"]]
    if late:
        owners = sorted({r["control_owner"] for r in late})
        actions.append(
            f"Chase {len(late)} overdue evidence request{'s' if len(late) != 1 else ''} "
            f"(owners: {', '.join(owners[:6])}{'…' if len(owners) > 6 else ''})."
        )
    drafts = requests["by_status_group"].get("draft", 0)
    if drafts:
        actions.append(f"Submit {drafts} drafted evidence request{'s' if drafts != 1 else ''} "
                       "for approval.")  # fmt: skip
    awaiting = status["review"]["controls_awaiting_review"]
    if awaiting:
        actions.append(f"Review {awaiting} control{'s' if awaiting != 1 else ''} whose testing "
                       "is signed off.")  # fmt: skip
    undecided = status["exceptions"]["failed_tests_without_deficiency"]
    if undecided:
        actions.append(
            f"Decide whether {undecided} failed test{'s' if undecided != 1 else ''} with no "
            "deficiency logged is a deficiency (auditor judgment)."
        )
    if status["exceptions"]["awaiting_decision"]:
        actions.append(
            f"Decide {status['exceptions']['awaiting_decision']} deficiency change"
            f"{'s' if status['exceptions']['awaiting_decision'] != 1 else ''} awaiting a decision."
        )
    for milestone in status["schedule"]["next_milestones"][:3]:
        actions.append(f'Meet the milestone "{milestone["milestone"]}" by {milestone["due"]}.')
    return actions or ["Nothing in the records needs action now."]


def _narrative(text: str) -> Section | None:
    cleaned = " ".join(text.split())
    if cleaned == "":
        return None
    return Section(NARRATIVE_LABEL, intro=NARRATIVE_NOTE, bullets=[cleaned], kind="narrative")


def build(
    connection: psycopg.Connection[Any],
    principal: Principal,
    organization_id: UUID,
    project_id: UUID,
    kind: str,
    filters: dict[str, Any],
    generated_by: str,
    narrative: str = "",
) -> Report:
    """Read the records and lay the report out. Every figure comes from engagement_data."""
    if kind not in REPORTS:
        raise ReportFailed(f"There is no report called {kind}.")

    def read(work: Callable[[], dict[str, Any]]) -> dict[str, Any]:
        try:
            return work()
        except data.DataUnavailable as error:
            raise ReportFailed(str(error)) from None

    sections: list[Section] = []
    landscape = kind != "executive_summary"
    if kind == "evidence_tracker":
        chosen = data.RequestFilters(**({"status_group": PENDING} | filters))
        result = read(
            lambda: data.evidence_requests(
                connection, principal, organization_id, project_id, chosen,
                sort="most_overdue", limit=None,
            )
        )  # fmt: skip
        scope = _scope_line(chosen.applied())
        sections.append(
            Section(
                "Summary",
                facts=[
                    ("Requests in this tracker", str(result["total"])),
                    ("Overdue", str(result["overdue"])),
                    ("Testing phase not recorded", str(result["phase_not_recorded"])),
                    *[
                        (data.GROUP_LABELS[key], str(count))
                        for key, count in result["by_status_group"].items()
                    ],
                ],
            )
        )
        sections.append(
            Section(
                "Requests",
                tables=[_request_table(result["requests"], result["total"], "Requests")],
            )
        )
        checks = result["data_quality"]
        rows = result["total"]
        primary = result
    elif kind == "control_summary":
        chosen_controls = data.ControlFilters(**filters)
        result = read(
            lambda: data.controls(
                connection, principal, organization_id, project_id, chosen_controls,
                limit=None,
            )
        )  # fmt: skip
        scope = _scope_line(chosen_controls.applied())
        sections.append(
            Section(
                "Summary",
                facts=[
                    ("Controls", str(result["total"])),
                    ("Failed testing", str(result["failed"])),
                    *[
                        (STATUS_WORDS.get(key, key), str(result["by_testing_status"].get(key, 0)))
                        for key in STATUS_WORDS
                    ],
                ],
            )
        )
        sections.append(
            Section("Controls", tables=[_control_table(result["controls"], result["total"])])
        )
        checks = result["data_quality"]
        rows = result["total"]
        primary = result
    elif kind == "exceptions_report":
        allowed = {key: filters[key] for key in ("include_closed", "application", "control_ref")
                   if key in filters}  # fmt: skip
        result = read(
            lambda: data.exceptions(
                connection, principal, organization_id, project_id, limit=None, **allowed
            )
        )
        scope = _scope_line(result["filters"]) if result["filters"] else "Open exceptions"
        sections.append(
            Section(
                "Summary",
                facts=[
                    ("Exceptions", str(result["total"])),
                    *[
                        (_words(status), str(count))
                        for status, count in result["by_status"].items()
                    ],
                ],
            )
        )
        sections.append(
            Section(
                "Exceptions",
                intro="Severity is shown only as the team recorded it. A failed test with no "
                "deficiency logged is listed for the auditor to judge; it is not a deficiency "
                "until one is logged and decided.",
                tables=[_exception_table(result["exceptions"], result["total"])],
            )
        )
        checks = []
        rows = result["total"]
        primary = result
    else:
        status = read(lambda: data.status(connection, principal, organization_id, project_id))
        requests = read(
            lambda: data.evidence_requests(
                connection, principal, organization_id, project_id,
                data.RequestFilters(status_group=PENDING), limit=None,
            )
        )  # fmt: skip
        scope = "The whole engagement"
        primary = status
        checks = status["data_quality"]
        rows = status["controls"]["in_scope"]
        controls = status["controls"]
        if kind == "status_report":
            exceptions = read(
                lambda: data.exceptions(
                    connection, principal, organization_id, project_id, limit=None
                )
            )
            review = read(
                lambda: data.review_items(
                    connection, principal, organization_id, project_id, limit=None
                )
            )
            late = [r for r in requests["requests"] if r["overdue"]]
            sections += [
                Section("Overall status", facts=_overview_facts(status)),
                Section(
                    "Control testing progress",
                    intro=controls["completion"]["rule"],
                    tables=[
                        Table(
                            "Controls by testing status",
                            ["Testing status", "Controls"],
                            [
                                [STATUS_WORDS[key], str(count)]
                                for key, count in controls["by_testing_status"].items()
                            ]
                            + [["Failed testing (any status)", str(controls["failed"])]],
                            widths=[3, 1],
                        )
                    ],
                ),
                Section(
                    "Evidence requests",
                    intro=status["evidence"]["rule"],
                    facts=[
                        (data.GROUP_LABELS[key], str(count))
                        for key, count in status["evidence"]["by_status_group"].items()
                    ],
                    tables=[
                        _request_table(
                            late,
                            status["evidence"]["overdue"],
                            "Overdue evidence requests",
                            empty="No evidence requests are overdue.",
                            titled=True,
                        )
                    ],
                ),
                Section(
                    "Open exceptions and deficiencies",
                    tables=[_exception_table(exceptions["exceptions"], exceptions["total"])],
                ),
                Section(
                    "Outstanding review items",
                    tables=[_review_table(review["items"], review["total"])],
                ),
                Section(
                    "Blockers and risks",
                    bullets=[
                        b["blocker"]
                        + (f": {b['oldest']}" if "oldest" in b else "")
                        + (f" ({', '.join(b['records'])})" if b.get("records") else "")
                        for b in status["blockers"]
                    ]
                    or ["None recorded."],
                ),
                Section("Next actions", bullets=_next_actions(status, requests)),
            ]
        else:
            done = controls["by_testing_status"].get("complete", 0)
            received = status["evidence"]["by_status_group"].get("received", 0)
            reviewing = controls["by_testing_status"].get("awaiting_review", 0)
            sections += [
                Section("Progress", facts=_overview_facts(status)[:8]),
                Section(
                    "Key accomplishments",
                    bullets=[
                        f"{done} of {controls['in_scope']} controls complete "
                        "(concluded, review signed off, no open review note).",
                        f"{received} evidence request{'s' if received != 1 else ''} received.",
                        f"{reviewing} control{'s' if reviewing != 1 else ''} tested and "
                        "awaiting review.",
                    ]
                    if done or received or reviewing
                    else [
                        "None yet: no control is complete, no evidence request has been "
                        "received, and no tested control awaits review."
                    ],
                ),
                Section(
                    "Significant outstanding items",
                    bullets=[
                        f"{status['evidence']['pending']} evidence requests pending, "
                        f"{status['evidence']['overdue']} overdue.",
                        f"{controls['failed']} controls failed testing.",
                        f"{status['exceptions']['open_deficiencies']} open deficiencies.",
                        f"{status['review']['open_review_notes']} open review notes.",
                    ],
                ),
                Section(
                    "Risks and blockers",
                    bullets=[
                        b["blocker"] + (f": {b['oldest']}" if "oldest" in b else "")
                        for b in status["blockers"]
                    ]
                    or ["None recorded."],
                ),
                Section(
                    "Decisions or escalations needed",
                    bullets=[
                        line
                        for line in _next_actions(status, requests)
                        if "Decide" in line or "Chase" in line
                    ]
                    or ["None recorded."],
                ),
                Section(
                    "Upcoming milestones",
                    bullets=[
                        f"{m['milestone']}: {m['due']} ({m['status'].replace('_', ' ').lower()})"
                        for m in status["schedule"]["next_milestones"]
                    ]
                    or ["None scheduled."],
                ),
                Section(
                    "Limitations",
                    bullets=[
                        f"The figures follow {status['rules']}.",
                        f"{len(status['data_quality'])} data check"
                        f"{'s' if len(status['data_quality']) != 1 else ''} found records "
                        "that contradict each other or miss values (see Data checks)."
                        if status["data_quality"]
                        else "The data checks found no records that contradict each other or "
                        "miss values.",
                        "This summary states no audit opinion: the only conclusions in it are "
                        "those recorded and signed off in the engagement.",
                    ],
                ),
            ]
    told = _narrative(narrative)
    if told is not None:
        sections.append(told)
    flagged = _checks(checks)
    if flagged is not None:
        sections.append(flagged)
    sections.append(_about(primary, scope, told is not None))
    report = Report(
        kind=kind,
        title=REPORTS[kind],
        engagement=str(primary["engagement"]),
        client=str(primary["client"]),
        as_of=str(primary["as_of"]),
        rules=str(primary["rules"]),
        scope=scope,
        generated_by=generated_by,
        generated_at=datetime.now(UTC),
        sections=sections,
        rows=rows,
        landscape=landscape,
    )
    check(report)
    return report


def check(report: Report) -> None:
    """A report must say what it covers and when, and hold every record it counted."""
    for name in ("engagement", "as_of", "rules", "scope", "title"):
        if not str(getattr(report, name)).strip():
            raise ReportFailed(f"The report is missing its {name.replace('_', ' ')}.")
    for section in report.sections:
        for table in section.tables:
            if table.total is not None and len(table.rows) != table.total:
                raise ReportFailed(
                    f"{table.title} holds {len(table.rows)} rows but {table.total} records "
                    "were counted. Nothing was saved; try again."
                )
            if any(len(row) != len(table.headers) for row in table.rows):
                raise ReportFailed(f"{table.title} has a row with missing columns.")


# Drawing


# Punctuation the PDF's fonts draw but readers cannot copy or search back: plain forms.
_PLAIN = str.maketrans({"·": "|", "—": "-", "–": "-", "…": "...", "“": '"', "”": '"',
                        "‘": "'", "’": "'", " ": " "})  # fmt: skip


def _records(count: int) -> str:
    return f"{count:,} record{'s' if count != 1 else ''}."


def _pdf_text(value: str) -> str:
    """Text for the PDF: markup escaped, punctuation plain, characters its fonts lack as ?."""
    safe = value.translate(_PLAIN).encode("latin-1", errors="replace").decode("latin-1")
    return escape(safe)


def render_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        CondPageBreak,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )
    from reportlab.platypus import (
        Table as PdfTable,
    )

    size = landscape(letter) if report.landscape else letter
    margin = 0.6 * inch
    out = io.BytesIO()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=9.5, alignment=TA_LEFT)
    head = ParagraphStyle("head", parent=small, textColor=colors.white, fontName="Helvetica-Bold")
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20, leading=24, alignment=TA_LEFT,
                        spaceAfter=4)  # fmt: skip
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13, leading=16,
                        spaceBefore=12, spaceAfter=5)  # fmt: skip
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=3)
    muted = ParagraphStyle("muted", parent=body, textColor=colors.HexColor("#4b5563"))
    story: list[Any] = [
        Paragraph(_pdf_text(report.title), h1),
        Paragraph(
            _pdf_text(
                f"{report.engagement}"
                + (f" · {report.client}" if report.client != "Not recorded" else "")
            ),
            muted,
        ),
        Paragraph(
            _pdf_text(
                f"Data as of {_moment(report.as_of)} · generated "
                f"{report.generated_at:%Y-%m-%d %H:%M} UTC by {report.generated_by} · "
                f"scope: {report.scope}"
            ),
            muted,
        ),
        Spacer(1, 8),
    ]
    width = size[0] - 2 * margin
    for section in report.sections:
        block: list[Any] = [Paragraph(_pdf_text(section.title), h2)]
        if section.intro:
            block.append(Paragraph(_pdf_text(section.intro), muted))
            block.append(Spacer(1, 4))
        if section.facts:
            facts = PdfTable(
                [[Paragraph(_pdf_text(k), body), Paragraph(f"<b>{_pdf_text(v)}</b>", body)]
                 for k, v in section.facts],
                colWidths=[width * 0.4, width * 0.6] if not report.landscape
                else [width * 0.3, width * 0.7],
            )  # fmt: skip
            facts.setStyle(
                TableStyle(
                    [
                        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ]
                )
            )
            block.append(facts)
        for line in section.bullets:
            block.append(Paragraph(_pdf_text(line), body, bulletText="•"))
        # A heading starts a new page rather than sit alone at the foot of one.
        story.append(CondPageBreak(1.3 * inch))
        story.append(KeepTogether(block) if not section.tables else block[0])
        if section.tables:
            story.extend(block[1:])
        if section.kind == "narrative":
            story.append(Spacer(1, 4))
        for table in section.tables:
            if table.titled:
                story.append(CondPageBreak(1.3 * inch))
                story.append(Paragraph(_pdf_text(table.title), h3))
            if not table.rows:
                story.append(Paragraph(_pdf_text(table.empty), muted))
                continue
            weights = table.widths or [1.0] * len(table.headers)
            scale = width / sum(weights)
            grid = PdfTable(
                [[Paragraph(_pdf_text(h), head) for h in table.headers]]
                + [[Paragraph(_pdf_text(cell), small) for cell in row] for row in table.rows],
                colWidths=[w * scale for w in weights],
                repeatRows=1,
            )
            grid.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                         [colors.white, colors.HexColor("#f6f7f9")]),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )  # fmt: skip
            story.append(grid)
            if table.total is not None:
                story.append(Paragraph(_pdf_text(_records(len(table.rows))), muted))

    def frame(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(
            margin,
            0.4 * inch,
            _pdf_text(f"{report.engagement} · {report.title} · data as of "
                      f"{_moment(report.as_of)} · {report.rules}").replace("&amp;", "&"),
        )  # fmt: skip
        canvas.drawRightString(size[0] - margin, 0.4 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        out,
        pagesize=size,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=0.7 * inch,
        title=f"{report.title}: {report.engagement}",
        author="AI Workforce",
    )
    document.build(story, onFirstPage=frame, onLaterPages=frame)
    return out.getvalue()


def render_docx(report: Report) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    if report.landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for edge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, edge, Inches(0.6))
    footer = section.footer.paragraphs[0]
    footer.text = (
        f"{report.engagement} · {report.title} · data as of {_moment(report.as_of)} · "
        f"{report.rules}"
    )
    for run in footer.runs:
        run.font.size = Pt(7.5)
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(report.title, level=0)
    line = document.add_paragraph()
    line.add_run(
        report.engagement + (f" · {report.client}" if report.client != "Not recorded" else "")
    ).bold = True
    meta = document.add_paragraph(
        f"Data as of {_moment(report.as_of)} · generated {report.generated_at:%Y-%m-%d %H:%M} "
        f"UTC by {report.generated_by} · scope: {report.scope}"
    )
    for run in meta.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    for part in report.sections:
        document.add_heading(part.title, level=1)
        if part.intro:
            intro = document.add_paragraph(part.intro)
            for run in intro.runs:
                run.italic = part.kind == "narrative"
                run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        if part.facts:
            facts = document.add_table(rows=0, cols=2)
            facts.style = "Table Grid"
            for label, value in part.facts:
                cells = facts.add_row().cells
                cells[0].text = label
                cells[1].text = value
                for run in cells[1].paragraphs[0].runs:
                    run.bold = True
        for text in part.bullets:
            document.add_paragraph(text, style="List Bullet")
        for table in part.tables:
            if table.titled:
                document.add_heading(table.title, level=2)
            if not table.rows:
                document.add_paragraph(table.empty)
                continue
            grid = document.add_table(rows=1, cols=len(table.headers))
            grid.style = "Table Grid"
            for cell, heading in zip(grid.rows[0].cells, table.headers, strict=True):
                cell.text = heading
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(8)
            for row in table.rows:
                cells = grid.add_row().cells
                for cell, value in zip(cells, row, strict=True):
                    cell.text = value
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8)
            if table.total is not None:
                document.add_paragraph(_records(len(table.rows)))
    out = io.BytesIO()
    document.core_properties.title = f"{report.title}: {report.engagement}"
    document.core_properties.author = "AI Workforce"
    document.save(out)
    return out.getvalue()
