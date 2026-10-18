"""Workpapers as .xlsx (Milestone 8.3): one workpaper, or the engagement's binder.

Each workpaper tab is a sheet (Design, Interim testing, Roll-forward): a title block
(reference, control, version, status, and sign-off), then each of the tab's sections in
order. Fields are a two-column table, narratives are wrapped text with their sources
listed under them, and tables keep their columns. The binder starts with an index sheet
that links to every sheet.

Cell values come from the records and from people. Anything that looks like a formula is
written as text, so a spreadsheet never runs it.
"""

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "1B2A41"
HEAD = PatternFill("solid", fgColor="E8EDF5")
TITLE = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="D6D9DF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WIDTH = 6


def _safe(value: Any) -> str:
    """Text for a cell. A leading = + - @ would start a formula, so it is quoted."""
    text = "" if value is None else str(value)
    text = text.replace("\x00", "")[:32_000]
    return f"'{text}" if text[:1] in ("=", "+", "-", "@") else text


def _status(body: Any) -> str:
    if body.status == "PREPARED":
        when = body.prepared_at.date().isoformat() if body.prepared_at else ""
        return f"Prepared by {body.prepared_by or ''} on {when} (version {body.prepared_version})"
    return "Draft"


def tabs_of(body: Any) -> list[tuple[Any, list[Any]]]:
    """A workpaper's tabs with their sections; a workpaper without tabs is one sheet."""
    content = body.content
    if not content.tabs:
        return [(None, list(content.sections))]
    return [
        (tab, [section for section in content.sections if section.tab == tab.key])
        for tab in content.tabs
    ]


def _sheet(sheet: Worksheet, engagement: str, body: Any, tab: Any, sections: list[Any]) -> None:
    content = body.content
    sheet.sheet_view.showGridLines = False
    widths = [24, 22, 22, 22, 22, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    row = 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=WIDTH)
    reference = content.reference if tab is None or not tab.reference else tab.reference
    cell = sheet.cell(row=row, column=1, value=_safe(f"{reference}  {content.title}"))
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = TITLE
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 26
    row += 1
    for label, value in (
        ("Engagement", engagement),
        ("Workpaper", content.template_title),
        ("Tab", "" if tab is None else tab.title),
        ("Purpose", content.purpose),
        ("Version", f"{body.shown_version} of {body.current_version}"),
        ("Status", _status(body)),
    ):
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, color="5A6475")
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=WIDTH)
        sheet.cell(row=row, column=2, value=_safe(value)).alignment = WRAP
        row += 1
    for section in sections:
        row += 1
        heading = sheet.cell(row=row, column=1, value=_safe(section.title))
        heading.font = Font(bold=True, size=12, color=NAVY)
        source = {"AGENT": "Drafted by the agent", "PERSON": "Written by a person"}.get(
            section.source, ""
        )
        if source:
            sheet.cell(row=row, column=WIDTH, value=source).font = Font(italic=True, color="8A93A3")
        row += 1
        if section.kind == "fields":
            for field in section.fields:
                label = sheet.cell(row=row, column=1, value=_safe(field.label))
                label.font = Font(bold=True)
                label.fill = HEAD
                label.border = BOX
                sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=WIDTH)
                value = sheet.cell(row=row, column=2, value=_safe(field.value))
                value.alignment = WRAP
                value.border = BOX
                row += 1
        elif section.kind == "text":
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=WIDTH)
            text = section.text or ""
            cell = sheet.cell(row=row, column=1, value=_safe(text))
            cell.alignment = WRAP
            sheet.row_dimensions[row].height = max(18, min(400, 15 * (len(text) // 110 + 1)))
            row += 1
            for citation in section.citations:
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=WIDTH)
                where = f", {citation.get('location')}" if citation.get("location") else ""
                note = sheet.cell(
                    row=row,
                    column=1,
                    value=_safe(
                        f"[{citation.get('number')}] {citation.get('file_name')}{where}: "
                        f"“{citation.get('quote')}”"
                    ),
                )
                note.font = Font(size=9, color="5A6475")
                note.alignment = WRAP
                row += 1
        else:
            columns = section.columns or [""]
            for index, name in enumerate(columns[:WIDTH], start=1):
                cell = sheet.cell(row=row, column=index, value=_safe(name))
                cell.font = Font(bold=True)
                cell.fill = HEAD
                cell.border = BOX
                cell.alignment = WRAP
            row += 1
            for values in section.rows:
                for index, value in enumerate(values[: len(columns[:WIDTH])], start=1):
                    cell = sheet.cell(row=row, column=index, value=_safe(value))
                    cell.border = BOX
                    cell.alignment = WRAP
                row += 1
    sheet.freeze_panes = "A2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1


def _title(body: Any, tab: Any, taken: set[str]) -> str:
    reference = body.content.reference if tab is None or not tab.reference else tab.reference
    base = "".join(ch for ch in reference if ch not in "[]:*?/\\")[:28] or "WP"
    title, n = base, 2
    while title in taken:
        title = f"{base[:26]}~{n}"
        n += 1
    taken.add(title)
    return title


def _bytes(book: Workbook) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def workpaper_xlsx(engagement: str, bodies: list[Any]) -> bytes:
    book = Workbook()
    first = book.active
    assert first is not None
    taken: set[str] = set()
    used_first = False
    for body in bodies:
        if body.content is None:
            continue
        for tab, sections in tabs_of(body):
            sheet = book.create_sheet() if used_first else first
            used_first = True
            sheet.title = _title(body, tab, taken)
            _sheet(sheet, engagement, body, tab, sections)
    return _bytes(book)


def binder_xlsx(engagement: str, bodies: list[Any]) -> bytes:
    """The engagement's workpapers in one workbook, with an index that links to each."""
    book = Workbook()
    index = book.active
    assert index is not None
    index.title = "Index"
    index.sheet_view.showGridLines = False
    for letter, width in zip("ABCDEF", (22, 44, 24, 12, 40, 10), strict=True):
        index.column_dimensions[letter].width = width
    index.merge_cells("A1:F1")
    head = index.cell(row=1, column=1, value=_safe(f"{engagement}: workpapers"))
    head.font = Font(bold=True, size=14, color="FFFFFF")
    head.fill = TITLE
    for column, name in enumerate(
        ("Reference", "Control", "Tab", "Version", "Status", "Sheet"), start=1
    ):
        cell = index.cell(row=3, column=column, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BOX
    taken = {"Index"}
    row = 4
    for body in bodies:
        if body.content is None:
            continue
        for tab, sections in tabs_of(body):
            title = _title(body, tab, taken)
            sheet = book.create_sheet(title)
            _sheet(sheet, engagement, body, tab, sections)
            values = (
                body.content.reference if tab is None else tab.reference,
                body.content.title,
                f"{body.content.template_title}: {tab.title}"
                if tab is not None
                else body.content.template_title,
                str(body.current_version),
                _status(body),
                title,
            )
            for column, value in enumerate(values, start=1):
                cell = index.cell(row=row, column=column, value=_safe(value))
                cell.border = BOX
                cell.alignment = WRAP
            link = index.cell(row=row, column=6)
            link.hyperlink = f"#'{title}'!A1"
            link.font = Font(color="1F5FD1", underline="single")
            row += 1
    if row == 4:
        index.cell(row=4, column=1, value="No workpapers have been drafted yet.")
    return _bytes(book)
