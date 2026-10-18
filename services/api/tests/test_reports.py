"""Engagement reports (Milestone 8.21): PDF and Word from the records, checked, labelled,
and offered like exports."""

import io
from datetime import UTC, datetime, timedelta
from typing import Any
from uuid import UUID

import pytest
from docx import Document
from pypdf import PdfReader
from support import Database
from test_console import _ask, _call, _reply, _Script, _thread, _turn
from test_engagement_data import _context, _engagement, _requests, _run, _set
from test_evidence_review import _base
from test_outbound import migrated

from ai_workforce_api import engagement_data, reports
from ai_workforce_api.console_tools import ToolFailed
from ai_workforce_api.db import connect

_ = migrated


def _pdf_text(content: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)


def _download(client: Any, team: dict[str, Any], export_id: str) -> Any:
    response = client.get(f"{_base(team['org'], team['project'])}/exports/{export_id}")
    assert response.status_code == 200, response.text
    return response


@pytest.mark.parametrize("report", list(reports.REPORTS))
def test_every_report_is_made_as_pdf_from_the_records(migrated: Database, report: str) -> None:
    team, client = _engagement(migrated, f"r-{report.split('_')[0]}")
    with connect(migrated.url) as connection:
        made = _run(_context(connection, migrated, team), "generate_report", report=report)
    file = made["export"]
    assert file["file_format"] == "pdf" and file["file_name"].endswith(".pdf")
    response = _download(client, team, file["id"])
    assert response.headers["content-type"] == "application/pdf"
    text = _pdf_text(response.content)
    # It says what it is, what it covers, when the records were read, and by which rules.
    assert reports.REPORTS[report] in text
    assert "SOX 2027" in text
    assert "Data as of" in text and file["as_of"][:10] in text
    assert "Metrics rulebook v1, pending approval" in text
    assert "About this report" in text
    assert "AI-written summary" not in text


def test_the_status_report_reads_cleanly(migrated: Database) -> None:
    team, client = _engagement(migrated, "rclean")
    with connect(migrated.url) as connection:
        made = _run(_context(connection, migrated, team), "generate_report", report="status_report")
    text = " ".join(_pdf_text(_download(client, team, made["export"]["id"]).content).split())
    moment = datetime.fromisoformat(made["export"]["as_of"]).astimezone(UTC)
    # The time as people read it, not a raw timestamp with its zone said twice.
    assert f"Data as of {moment:%Y-%m-%d %H:%M:%S} UTC" in text
    assert "+00:00" not in text
    # An empty table says which one is empty.
    assert "Overdue evidence requests No evidence requests are overdue." in text
    # A breakdown is not a list of records, so it is not counted as one.
    assert "6 records." not in text
    assert "This report states no audit opinion: the only conclusions in it are" in text


def test_long_text_is_cut_at_a_word_and_marked() -> None:
    assert reports._clip("short", 10) == "short"
    assert reports._clip("We need your written procedure today", 20) == "We need your written…"
    assert reports._words("NOT_CONCLUDED") == "Not concluded"


def test_the_executive_summary_states_the_figures_and_its_limits(migrated: Database) -> None:
    team, client = _engagement(migrated, "exec")
    now = datetime.now(UTC)
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(connection, ids[0], status="SENT", due_at=now - timedelta(days=6))
        connection.commit()
        context = _context(connection, migrated, team)
        status = _run(context, "engagement_status")
        made = _run(context, "generate_report", report="executive_summary")
    text = " ".join(_pdf_text(_download(client, team, made["export"]["id"]).content).split())
    controls = status["controls"]
    assert f"Controls in scope {controls['in_scope']}" in text
    assert f"{status['evidence']['pending']} evidence requests pending" in text
    assert f"{status['evidence']['overdue']} overdue" in text
    for heading in (
        "Key accomplishments",
        "Significant outstanding items",
        "Risks and blockers",
        "Decisions or escalations needed",
        "Upcoming milestones",
        "Limitations",
    ):
        assert heading in text
    assert "states no audit opinion" in text


def test_the_evidence_tracker_holds_exactly_the_pending_requests(migrated: Database) -> None:
    team, client = _engagement(migrated, "tracker")
    now = datetime.now(UTC)
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(
            connection,
            ids[0],
            status="SENT",
            due_at=now - timedelta(days=3),
            testing_phase="INTERIM",
        )
        connection.commit()
        context = _context(connection, migrated, team)
        pending = engagement_data.evidence_requests(
            connection, context.principal, team["org"], team["project"],
            engagement_data.RequestFilters(status_group=reports.PENDING), limit=None,
        )  # fmt: skip
        interim = _run(
            context,
            "generate_report",
            report="evidence_tracker",
            file_format="docx",
            filters={"phase": "INTERIM"},
        )
        everything = _run(context, "generate_report", report="evidence_tracker")
    assert everything["export"]["rows"] == pending["total"]
    document = Document(io.BytesIO(_download(client, team, interim["export"]["id"]).content))
    tables = document.tables
    grid = tables[-1]
    headers = [cell.text for cell in grid.rows[0].cells]
    assert headers[:4] == ["Request", "Control", "Application", "Phase"]
    rows = [[cell.text for cell in row.cells] for row in grid.rows[1:]]
    expected = sum(1 for r in pending["requests"] if r["phase"] == "INTERIM")
    assert len(rows) == interim["export"]["rows"] == expected
    assert all(row[3] == "Interim" for row in rows)
    late = next(row for row in rows if row[0].startswith(pending["requests"][0]["reference"]))
    assert late[headers.index("Days overdue")] == "3"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "scope: status group: draft or awaiting approval or with client or partly " in text
    assert "phase: INTERIM" in text


def test_an_ai_narrative_is_labelled_and_kept_with_the_file(migrated: Database) -> None:
    team, client = _engagement(migrated, "narrative")
    with connect(migrated.url) as connection:
        made = _run(
            _context(connection, migrated, team),
            "generate_report",
            report="status_report",
            narrative="Evidence is the main risk: most requests are still drafts.",
        )
        record = connection.execute(
            "SELECT narrative, filters FROM engagement_exports WHERE id = %s",
            (made["export"]["id"],),
        ).fetchone()
        audit = connection.execute(
            "SELECT payload FROM audit_events WHERE target_id = %s AND action = 'export.created'",
            (made["export"]["id"],),
        ).fetchone()
    text = " ".join(_pdf_text(_download(client, team, made["export"]["id"]).content).split())
    label = text.index("AI-written summary")
    assert "not verified, and not an audit conclusion" in text[label:]
    assert "Evidence is the main risk" in text[label:]
    # The figures come before it, apart from it.
    assert text.index("Overall status") < label
    assert 'The section "AI-written summary" is the only AI-written part.' in text
    assert record is not None and record[0].startswith("Evidence is the main risk")
    assert record[1] == {"scope": "The whole engagement"}
    assert audit is not None and audit[0]["narrative"] is True


def test_a_report_that_does_not_add_up_is_never_saved(
    migrated: Database, monkeypatch: pytest.MonkeyPatch
) -> None:
    team, _client = _engagement(migrated, "addsup")
    original = engagement_data.evidence_requests

    def short(*args: Any, **kwargs: Any) -> dict[str, Any]:
        result = original(*args, **kwargs)
        result["requests"] = result["requests"][:-1]
        return result

    monkeypatch.setattr(engagement_data, "evidence_requests", short)
    with connect(migrated.url) as connection:
        with pytest.raises(ToolFailed, match="rows but .* records were counted"):
            _run(_context(connection, migrated, team), "generate_report", report="evidence_tracker")
        saved = connection.execute(
            "SELECT count(*) FROM engagement_exports WHERE project_id = %s", (team["project"],)
        ).fetchone()
    assert saved == (0,)


def test_reports_and_records_each_take_their_own_formats(migrated: Database) -> None:
    team, _client = _engagement(migrated, "formats")
    with connect(migrated.url) as connection:
        context = _context(connection, migrated, team)
        with pytest.raises(ToolFailed, match="not valid"):
            _run(context, "generate_report", report="status_report", file_format="xlsx")
        with pytest.raises(ToolFailed, match="not valid"):
            _run(context, "export_records", dataset="controls", file_format="pdf")
        with pytest.raises(ToolFailed, match="takes no filters"):
            _run(context, "generate_report", report="status_report", filters={"phase": "INTERIM"})


def test_markup_in_the_records_cannot_break_a_report(migrated: Database) -> None:
    team, client = _engagement(migrated, "markup")
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(
            connection,
            ids[0],
            evidence_description="<b>Bold</b> & <script>x</script> ✓ 合同 list of users",
        )
        connection.commit()
        made = _run(
            _context(connection, migrated, team), "generate_report", report="evidence_tracker"
        )
    text = " ".join(_pdf_text(_download(client, team, made["export"]["id"]).content).split())
    # Shown as written, never read as markup; what the fonts cannot draw shows as ?.
    assert "<b>Bold</b> & <script>x</script> ? ?? list of users" in text


def test_a_report_is_only_for_people_who_can_see_the_engagement(migrated: Database) -> None:
    team, _client = _engagement(migrated, "rpaccess")
    _other, other_client = _engagement(migrated, "rpother")
    with connect(migrated.url) as connection:
        made = _run(
            _context(connection, migrated, team), "generate_report", report="control_summary"
        )
    path = f"{_base(team['org'], team['project'])}/exports/{made['export']['id']}"
    assert other_client.get(path).status_code == 404


def test_the_console_offers_the_pdf_with_the_answer(migrated: Database) -> None:
    team, client = _engagement(migrated, "consolepdf")
    message_id = _ask(client, team, "Can I have that as a PDF?")
    script = _Script(
        [
            _reply("", _call("generate_report", report="status_report")),
            _reply("Here is the status report as a PDF."),
        ]
    )
    with connect(migrated.url) as connection:
        assert _turn(connection, migrated, script, message_id) == "DONE"
    step = _thread(client, team)["messages"][-1]["actions"][0]
    assert step["status"] == "DONE"
    assert step["summary"] == "Made the engagement status report as PDF"
    assert step["result"]["export"]["file_format"] == "pdf"
    assert UUID(step["result"]["export"]["id"])
