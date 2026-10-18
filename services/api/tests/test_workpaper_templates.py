"""A firm's own workpaper templates: read from Word or Excel into the workbook's tabs,
customised from a standard one in the editor, offered next to the standard ones for
their kind of control, and followed by previews and workpapers."""

import io
import zipfile
from typing import Any

import openpyxl
from support import Database
from test_onboarding import _base, _new_engagement
from test_outbound import migrated

from ai_workforce_api.workpaper_templates import (
    TemplateFileError,
    parse_docx,
    parse_xlsx,
    period_of,
    sections_from,
    tabs_from,
)
from ai_workforce_api.workpapers import TEMPLATES, default_template, tabs_for

__all__ = ["migrated"]

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx(blocks: list[tuple[str, str | list[list[str]]]]) -> bytes:
    """A minimal Word document: headings, paragraphs, and tables, in order."""
    parts: list[str] = []
    for kind, value in blocks:
        if isinstance(value, str):
            if kind == "heading":
                style = '<w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
                parts.append(f"<w:p>{style}<w:r><w:t>{value}</w:t></w:r></w:p>")
            else:
                parts.append(f"<w:p><w:r><w:t>{value}</w:t></w:r></w:p>")
        else:
            rows = "".join(
                "<w:tr>"
                + "".join(f"<w:tc><w:p><w:r><w:t>{cell}</w:t></w:r></w:p></w:tc>" for cell in row)
                + "</w:tr>"
                for row in value
            )
            parts.append(f"<w:tbl>{rows}</w:tbl>")
    document = (
        f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{_W}"><w:body>'
        + "".join(parts)
        + "</w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            'content-types"><Default Extension="xml" ContentType="application/xml"/></Types>',
        )
        archive.writestr("word/document.xml", document)
    return buffer.getvalue()


# A firm's testing workpaper: no tab headings, so it covers one part of the workbook.
FIRM_DOCX = _docx(
    [
        ("heading", "Control information"),
        (
            "table",
            [
                ["Control reference", ""],
                ["Control owner", ""],
                ["Risk addressed", ""],
                ["Period", ""],
            ],
        ),
        ("heading", "Understanding of the process"),
        ("text", "Describe the process from initiation to recording, naming the systems used."),
        ("heading", "Test procedures and results"),
        ("table", [["Sample", "Attribute 1", "Attribute 2", "Result"]]),
        ("heading", "Conclusion"),
        ("text", "State whether the control operated effectively."),
        ("heading", "Sign-off"),
        ("table", [["Prepared by", "", ""], ["Reviewed by", "", ""]]),
    ]
)

# A firm's whole workpaper: top headings name the tabs.
FIRM_WORKBOOK_DOCX = _docx(
    [
        ("heading", "Design and walkthrough"),
        ("heading", "Process description"),
        ("text", "Describe the process."),
        ("heading", "Walkthrough"),
        ("table", [["Step", "Evidence", "Result"]]),
        ("heading", "Interim testing"),
        ("heading", "Testing procedures"),
        ("table", [["Sample", "Attribute", "Result"]]),
        ("heading", "Conclusion"),
        ("text", "Conclude on the interim period."),
    ]
)


def test_a_word_template_is_read_for_its_sections() -> None:
    parts = parse_docx(FIRM_DOCX)
    assert [name for name, _blocks in parts] == [""]
    sections = sections_from(parts[0][1])
    assert [(s.title, s.kind) for s in sections] == [
        ("Control information", "fields"),
        ("Understanding of the process", "guidance"),
        ("Test procedures and results", "matrix"),
        ("Conclusion", "guidance"),
        ("Sign-off", "signoff"),
    ]
    assert sections[1].guidance.startswith("Describe the process")
    assert sections[2].columns == ("Sample", "Attribute 1", "Attribute 2", "Result")


def test_tab_headings_split_a_word_template_into_the_workbooks_tabs() -> None:
    parts = parse_docx(FIRM_WORKBOOK_DOCX)
    assert [name for name, _blocks in parts] == ["Design and walkthrough", "Interim testing"]
    tabs = tabs_from(parts, "whole", "BUSINESS")
    assert [(tab.key, tab.period, tab.title) for tab in tabs] == [
        ("design", "design", "Design and walkthrough"),
        ("interim", "interim", "Interim testing"),
        ("rollforward", "rollforward", "Roll-forward"),
    ]
    assert [s.title for s in tabs[0].sections] == ["Process description", "Walkthrough"]
    assert [s.title for s in tabs[1].sections] == ["Testing procedures", "Conclusion"]
    # The roll-forward tab the file lacks is the standard one for the kind.
    assert tabs[2].sections[1].title == "Changes since interim"
    assert period_of("RFP") == "rollforward" and period_of("Instructions") is None
    # A section heading inside a document is not a tab; a sheet named Testing is.
    assert period_of("Understanding of the process") is None
    assert period_of("Testing", sheet=True) == "interim"


def test_a_file_that_covers_one_tab_takes_that_tab() -> None:
    tabs = tabs_from(parse_docx(FIRM_DOCX), "interim", "ITGC")
    assert [(tab.key, tab.title) for tab in tabs] == [
        ("design", "Design"),
        ("interim", "Interim testing"),
        ("rollforward", "Roll-forward"),
    ]
    assert tabs[1].sections[0].title == "Control information"
    assert tabs[0].sections[1].title == "Understanding of the control"
    try:
        tabs_from(parse_docx(FIRM_DOCX), "whole", "ITGC")
    except TemplateFileError as error:
        assert "Design, Interim, or Roll-forward" in error.detail
    else:
        raise AssertionError("expected a refusal")


def test_an_excel_workbook_is_read_sheet_by_sheet() -> None:
    workbook = openpyxl.Workbook()
    design = workbook.active
    assert design is not None
    design.title = "Design"
    for row in [
        ["Walkthrough"],
        ["Follow one instance of the control and note each step."],
        ["Step", "Performed by", "Evidence", "Observation"],
        ["1", "", "", ""],
        [],
        ["Prepared by and reviewed by"],
        ["Role", "Name", "Date"],
    ]:
        design.append(row)
    interim = workbook.create_sheet("Interim TOE")
    for row in [
        ["Checks to perform"],
        ["Check", "Done", "Result"],
        ["Agree the listing to the system"],
        ["Check approvals"],
    ]:
        interim.append(row)
    workbook.create_sheet("Notes").append(["Internal notes"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    parts = parse_xlsx(buffer.getvalue())
    assert [name for name, _blocks in parts] == [
        "sheet:Design",
        "sheet:Interim TOE",
        "sheet:Notes",
    ]
    sections = sections_from(parts[0][1])
    assert [(s.title, s.kind) for s in sections] == [
        ("Walkthrough", "table"),
        ("Prepared by and reviewed by", "signoff"),
    ]
    assert sections[0].columns == ("Step", "Performed by", "Evidence", "Observation")
    assert sections[0].guidance.startswith("Follow one instance")
    tabs = tabs_from(parts, "whole", "KEY_REPORT")
    assert [tab.period for tab in tabs] == ["design", "interim", "rollforward"]
    # Rows with only their first cell filled are the firm's starter rows.
    procedures = tabs[1].sections[0]
    assert procedures.columns == ("Check", "Done", "Result")
    assert procedures.rows == (("Agree the listing to the system",), ("Check approvals",))


def test_a_file_without_structure_is_refused() -> None:
    try:
        sections_from([])
    except TemplateFileError as error:
        assert "no headings" in error.detail
    else:
        raise AssertionError("expected a refusal")
    try:
        parse_docx(b"not a zip")
    except TemplateFileError as error:
        assert "Word" in error.detail
    else:
        raise AssertionError("expected a refusal")


def test_each_kind_of_control_has_a_standard_workbook() -> None:
    assert [t.key for t in TEMPLATES] == ["itgc", "itac", "business", "key-report"]
    for template in TEMPLATES:
        assert [tab.period for tab in template.tabs] == ["design", "interim", "rollforward"]
        assert [tab.key for tab in tabs_for(template, "DESIGN")] == ["design"]
        assert [tab.key for tab in tabs_for(template, "EFFECTIVENESS")] == [
            "interim",
            "rollforward",
        ]
        for tab in template.tabs:
            assert tab.sections[0].kind == "fields"
            assert tab.sections[-1].kind == "signoff"
    assert default_template("ACCESS") == "itgc"
    assert default_template("ITAC") == "itac"
    assert default_template("IPE") == "key-report"
    assert default_template("ORDER_TO_CASH") == "business"


def _control_body(control: dict[str, Any], template_key: str) -> dict[str, Any]:
    return {
        "version": control["version"],
        "included": True,
        "key_control": True,
        "owner_name": "",
        "owner_email": "",
        "frequency": control["frequency"],
        "nature": control["nature"],
        "application_ids": control["application_ids"],
        "template_key": template_key,
    }


def test_a_firms_template_is_offered_applied_previewed_and_removed(migrated: Database) -> None:
    team = _new_engagement(migrated, "firm-templates")
    owner, base = team["owner"], _base(team["org"], team["project"])
    owner.post(f"{base}/applications", json={"name": "SAP ERP"})
    owner.post(f"{base}/controls/load-library", json={"domains": ["ACCESS", "ITAC"]})
    assert owner.post(f"{base}/templates/confirm").status_code == 200
    uploaded = owner.post(
        f"{base}/workpaper-templates",
        params={"file_name": "Firm OE workpaper.docx", "kind": "ITGC", "part": "interim"},
        content=FIRM_DOCX,
    )
    assert uploaded.status_code == 200, uploaded.text
    body = uploaded.json()
    firm = [t for t in body["templates"] if t["source"] == "FIRM"]
    assert len(firm) == 1
    template = firm[0]
    assert template["kind"] == "ITGC"
    assert template["file_name"] == "Firm OE workpaper.docx"
    assert [tab["period"] for tab in template["tabs"]] == ["design", "interim", "rollforward"]
    assert template["tabs"][1]["sections"][0]["title"] == "Control information"
    # A new template means the team confirms the templates again.
    assert body["settings"]["templates_confirmed_at"] is None
    # Nothing uses it until the team applies it; then every ITGC follows it and the
    # application controls keep theirs.
    assert all(c["template_key"] != template["key"] for c in body["controls"])
    applied = owner.post(f"{base}/workpaper-templates/{template['template_id']}/apply")
    assert applied.status_code == 200, applied.text
    controls = applied.json()["controls"]
    itgcs = [c for c in controls if c["domain"] == "ACCESS"]
    assert itgcs and all(c["template_key"] == template["key"] for c in itgcs)
    assert all(c["template_key"] == "itac" for c in controls if c["domain"] == "ITAC")
    # A preview fills the firm's tabs with a real control.
    control = itgcs[0]
    shown = owner.get(
        f"{base}/workpaper-preview",
        params={"template": template["key"], "control_id": control["id"]},
    )
    assert shown.status_code == 200, shown.text
    assert shown.json()["template_title"] == "Firm OE workpaper"
    tabs = {tab["key"]: tab for tab in shown.json()["tabs"]}
    assert tabs["interim"]["reference"].endswith("-TOE")
    sections = {s["title"]: s for s in tabs["interim"]["sections"]}
    assert any(f["label"] == "Control" for f in sections["Control information"]["fields"])
    assert sections["Test procedures and results"]["columns"][0] == "Sample"
    assert sections["Sign-off"]["rows"][0][0] == "Prepared by"
    # An application control cannot take an ITGC template: it keeps its own kind's.
    itac = next(c for c in controls if c["domain"] == "ITAC")
    kept = owner.patch(f"{base}/controls/{itac['id']}", json=_control_body(itac, template["key"]))
    assert kept.status_code == 200, kept.text
    assert next(c for c in kept.json()["controls"] if c["id"] == itac["id"])["template_key"] == (
        "itac"
    )
    # Removing it puts the standard template back.
    removed = owner.post(f"{base}/workpaper-templates/{template['template_id']}/remove")
    assert removed.status_code == 200, removed.text
    assert all(
        c["template_key"] == "itgc" for c in removed.json()["controls"] if c["domain"] == "ACCESS"
    )
    assert not any(t["source"] == "FIRM" for t in removed.json()["templates"])
    refused = owner.post(
        f"{base}/workpaper-templates",
        params={"file_name": "notes.txt", "kind": "ITGC"},
        content=b"plain text",
    )
    assert refused.status_code == 422
    unnamed = owner.post(
        f"{base}/workpaper-templates",
        params={"file_name": "Firm OE workpaper.docx", "kind": "ITGC"},
        content=FIRM_DOCX,
    )
    assert unnamed.status_code == 422
    assert "Design, Interim, or Roll-forward" in unnamed.json()["detail"]


def test_a_standard_template_is_customised_and_saved_as_the_firms_own(
    migrated: Database,
) -> None:
    team = _new_engagement(migrated, "custom-templates")
    owner, base = team["owner"], _base(team["org"], team["project"])
    owner.post(f"{base}/applications", json={"name": "SAP ERP"})
    owner.post(f"{base}/controls/load-library", json={"domains": ["ORDER_TO_CASH"]})
    made = owner.post(f"{base}/workpaper-templates/customize", json={"base_key": "business"})
    assert made.status_code == 200, made.text
    template = next(t for t in made.json()["templates"] if t["source"] == "FIRM")
    assert template["title"] == "Business process control (your firm)"
    assert template["base_key"] == "business"
    assert template["kind"] == "BUSINESS"
    assert [tab["title"] for tab in template["tabs"]] == [
        "Design",
        "Interim testing",
        "Roll-forward",
    ]
    # The firm renames a tab, drops the segregation of duties section, adds its own
    # procedure table, and saves.
    tabs = template["tabs"]
    design = tabs[0]
    design["title"] = "Design and implementation"
    design["sections"] = [
        s for s in design["sections"] if s["title"] != "Segregation of duties"
    ] + [
        {
            "title": "Firm procedures",
            "kind": "table",
            "guidance": "Perform each procedure.",
            "columns": ["Procedure", "Result"],
            "rows": [["Inspect the approval", ""]],
        }
    ]
    saved = owner.put(
        f"{base}/workpaper-templates/{template['template_id']}",
        json={
            "title": "Our BP workpaper",
            "version": template["version"],
            "tabs": [
                {"title": tab["title"], "period": tab["period"], "sections": tab["sections"]}
                for tab in tabs
            ],
        },
    )
    assert saved.status_code == 200, saved.text
    edited = next(t for t in saved.json()["templates"] if t["source"] == "FIRM")
    assert edited["title"] == "Our BP workpaper"
    assert edited["version"] == template["version"] + 1
    assert edited["tabs"][0]["title"] == "Design and implementation"
    titles = [s["title"] for s in edited["tabs"][0]["sections"]]
    assert "Segregation of duties" not in titles and titles[-1] == "Firm procedures"
    # A stale save is refused; a bad body too.
    stale = owner.put(
        f"{base}/workpaper-templates/{template['template_id']}",
        json={"title": "Again", "version": template["version"], "tabs": edited["tabs"]},
    )
    assert stale.status_code == 409
    bad = owner.put(
        f"{base}/workpaper-templates/{template['template_id']}",
        json={"title": "", "version": edited["version"], "tabs": []},
    )
    assert bad.status_code == 422
    # Applied, the business process controls follow it and the preview shows the firm's
    # procedure table with its starter row.
    applied = owner.post(f"{base}/workpaper-templates/{template['template_id']}/apply")
    assert applied.status_code == 200
    control = next(c for c in applied.json()["controls"] if c["template_key"] == edited["key"])
    shown = owner.get(
        f"{base}/workpaper-preview",
        params={"template": edited["key"], "control_id": control["id"]},
    ).json()
    design_tab = shown["tabs"][0]
    assert design_tab["title"] == "Design and implementation"
    procedures = next(s for s in design_tab["sections"] if s["title"] == "Firm procedures")
    assert procedures["rows"] == [["Inspect the approval", ""]]
