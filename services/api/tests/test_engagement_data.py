"""Engagement questions and exports (Milestone 8.20): answers from the records, totals the
application works out, exports that match them, and nothing read beyond what the person
may see or changed by reading."""

import io
from datetime import UTC, datetime, timedelta
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from support import Database
from test_console import _ask, _call, _reply, _Script, _thread, _turn
from test_evidence_review import _base, _client, _sent_engagement_request
from test_execution import _add_project_member, _add_user
from test_outbound import migrated

from ai_workforce_api import engagement_data, exports
from ai_workforce_api.authorization import Principal, PrincipalKind
from ai_workforce_api.console_tools import ToolContext, ToolFailed, execute, parse
from ai_workforce_api.db import connect, independent_steps
from ai_workforce_api.settings import Settings
from ai_workforce_api.storage import StoreError

_ = migrated


def _context(connection: Any, database: Database, team: dict[str, Any]) -> ToolContext:
    # As the console runs a turn: each read is its own transaction.
    independent_steps(connection)
    return ToolContext(
        connection=connection,
        settings=Settings(database_url=database.url),
        principal=Principal(kind=PrincipalKind.USER, id=team["owner_id"]),
        organization_id=team["org"],
        project_id=team["project"],
    )


def _run(context: ToolContext, name: str, **arguments: Any) -> dict[str, Any]:
    tool, parsed = parse(name, arguments)
    return execute(tool, parsed, context)


def _requests(connection: Any, project: UUID) -> list[UUID]:
    return [
        UUID(str(row[0]))
        for row in connection.execute(
            "SELECT id FROM pbc_requests WHERE project_id = %s ORDER BY control_ref, created_at",
            (project,),
        ).fetchall()
    ]


def _set(connection: Any, request_id: UUID, **columns: Any) -> None:
    status = columns.get("status")
    # Anything past approval has gone out in an email thread (pbc_requests_message_required).
    if status is not None and status not in ("DRAFT", "APPROVAL_PENDING", "CANCELLED", "CLOSED"):
        columns.setdefault("outbound_message_id", f"msg-{uuid4().hex[:8]}")
        columns.setdefault("thread_id", f"thread-{uuid4().hex[:8]}")
    assignments = ", ".join(f"{name} = %s" for name in columns)
    connection.execute(
        f"UPDATE pbc_requests SET {assignments} WHERE id = %s", (*columns.values(), request_id)
    )


def _engagement(database: Database, slug: str) -> tuple[dict[str, Any], Any]:
    team, client, _request = _sent_engagement_request(database, slug, "ITGC-AC-01")
    return team, client


def test_the_status_is_worked_out_from_the_records(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "status")
    now = datetime.now(UTC)
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(connection, ids[0], status="SENT", due_at=now - timedelta(days=3))
        _set(connection, ids[1], status="RECEIVED", received_at=now)
        connection.commit()
        context = _context(connection, migrated, team)
        result = _run(context, "engagement_status")
        expected = connection.execute(
            """
            SELECT (SELECT count(*) FROM engagement_controls WHERE project_id = %s AND included),
                   (SELECT count(*) FROM pbc_requests WHERE project_id = %s
                    AND status NOT IN ('CANCELLED', 'RECEIVED', 'CLOSED')),
                   (SELECT count(*) FROM pbc_requests WHERE project_id = %s
                    AND status IN ('SENT', 'WAITING_FOR_RESPONSE', 'PARTIALLY_RECEIVED',
                                   'OVERDUE', 'ESCALATED') AND due_at < now())
            """,
            (team["project"], team["project"], team["project"]),
        ).fetchone()
    assert expected is not None
    assert result["controls"]["in_scope"] == expected[0] > 0
    assert result["evidence"]["pending"] == expected[1]
    assert result["evidence"]["overdue"] == expected[2] >= 1
    completion = result["controls"]["completion"]
    assert completion["in_scope"] == expected[0] and completion["complete"] == 0
    assert "review signed off" in completion["rule"]
    assert result["rules"] == "Metrics rulebook v1, pending approval"
    assert datetime.fromisoformat(result["as_of"]) <= datetime.now(UTC)
    late = next(b for b in result["blockers"] if "overdue" in b["blocker"])
    assert any(ref.startswith("PBC-") for ref in late["records"])
    # Received with nothing linked: reported as it is, not corrected.
    checks = {check["check"] for check in result["data_quality"]}
    assert "received_without_files" in checks


def test_overdue_follows_the_rule_at_its_edges(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "overdue")
    now = datetime.now(UTC)
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        # Start from nothing sent, so only the cases below can be late.
        connection.execute(
            "UPDATE pbc_requests SET status = 'DRAFT', outbound_message_id = NULL, "
            "thread_id = NULL, received_at = NULL WHERE project_id = %s",
            (team["project"],),
        )
        cases = {
            "yesterday": (ids[0], {"status": "SENT", "due_at": now - timedelta(days=1)}),
            "earlier_today": (ids[1], {"status": "SENT", "due_at": now - timedelta(minutes=5)}),
            "tomorrow": (ids[2], {"status": "SENT", "due_at": now + timedelta(days=1)}),
            "draft_late": (ids[3], {"status": "DRAFT", "due_at": now - timedelta(days=7)}),
            "cancelled": (
                ids[4],
                {
                    "status": "CANCELLED",
                    "due_at": now - timedelta(days=7),
                    "cancellation_reason": "Not needed",
                },
            ),
            "partial": (
                ids[5],
                {"status": "PARTIALLY_RECEIVED", "due_at": now - timedelta(days=2)},
            ),
            "received": (
                ids[6],
                {"status": "RECEIVED", "due_at": now - timedelta(days=2), "received_at": now},
            ),
        }
        for request_id, columns in cases.values():
            _set(connection, request_id, **columns)
        connection.commit()
        result = _run(_context(connection, migrated, team), "list_requests", overdue_only=True)
    late = {UUID(row["request_id"]): row["days_overdue"] for row in result["requests"]}
    assert late == {
        cases["yesterday"][0]: 1,
        cases["partial"][0]: 2,
        cases["earlier_today"][0]: 0,
    }
    assert result["total"] == 3 and result["overdue"] == 3
    # Sorted most overdue first.
    assert [row["days_overdue"] for row in result["requests"]] == [2, 1, 0]


def test_pending_evidence_by_phase_and_phases_not_recorded(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "phase")
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(connection, ids[0], testing_phase=None)
        _set(connection, ids[1], testing_phase=None)
        connection.commit()
        context = _context(connection, migrated, team)
        pending = ["draft", "awaiting_approval", "with_client", "partly_received"]
        interim = _run(context, "list_requests", status_group=pending, phase="INTERIM")
        unknown = _run(context, "list_requests", status_group=pending, phase="NOT_RECORDED")
        everything = _run(context, "list_requests", status_group=pending)
        expected = connection.execute(
            "SELECT count(*) FILTER (WHERE testing_phase = 'INTERIM'), "
            "count(*) FILTER (WHERE testing_phase IS NULL), count(*) FROM pbc_requests "
            "WHERE project_id = %s AND status NOT IN ('CANCELLED', 'RECEIVED', 'CLOSED')",
            (team["project"],),
        ).fetchone()
    assert expected is not None
    assert interim["total"] == expected[0]
    assert all(row["phase"] == "INTERIM" for row in interim["requests"])
    # A phase the record does not hold is said, never dropped or counted as zero.
    assert unknown["total"] == expected[1] == 2
    assert everything["phase_not_recorded"] == 2
    assert everything["total"] == expected[2]
    assert any(c["check"] == "pending_without_phase" for c in everything["data_quality"])


def test_filters_by_application_owner_and_process(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "filters")
    with connect(migrated.url) as connection:
        connection.execute(
            "UPDATE engagement_controls SET owner_name = 'Dana Whitfield' "
            "WHERE project_id = %s AND control_ref = 'ITGC-AC-02'",
            (team["project"],),
        )
        connection.commit()
        context = _context(connection, migrated, team)
        # Every record, not one page of them, to check the filtered totals against.
        everything = engagement_data.evidence_requests(
            connection,
            context.principal,
            team["org"],
            team["project"],
            engagement_data.RequestFilters(),
            limit=None,
        )
        dana = _run(context, "list_requests", owner="dana", limit=100)
        application = everything["requests"][0]["application"]
        by_app = _run(context, "list_requests", application=application, limit=100)
        process = everything["requests"][0]["process"]
        by_process = _run(context, "list_requests", process=process, limit=100)
        controls = _run(context, "list_controls", owner="Dana Whitfield")
    assert dana["total"] > 0
    assert all(row["control_ref"].startswith("ITGC-AC-02") for row in dana["requests"])
    assert dana["filters"] == {"owner": "dana"}
    assert by_app["total"] == sum(
        1 for r in everything["requests"] if r["application"] == application
    )
    assert by_process["total"] == sum(1 for r in everything["requests"] if r["process"] == process)
    assert controls["total"] >= 1
    assert all(row["owner"] == "Dana Whitfield" for row in controls["controls"])


def test_an_excel_export_holds_exactly_the_filtered_records(migrated: Database) -> None:
    team, client = _engagement(migrated, "excel")
    now = datetime.now(UTC)
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(connection, ids[0], status="SENT", due_at=now - timedelta(days=4))
        # A description that would be a formula if it were not kept as text.
        _set(connection, ids[0], evidence_description='=HYPERLINK("http://x.example","go")')
        connection.commit()
        context = _context(connection, migrated, team)
        shown = _run(context, "list_requests", overdue_only=True)
        made = _run(
            context,
            "export_records",
            dataset="evidence_requests",
            file_format="xlsx",
            filters={"overdue_only": True},
        )
        record = connection.execute(
            "SELECT row_count, as_of, filters, expires_at - created_at FROM engagement_exports "
            "WHERE id = %s",
            (made["export"]["id"],),
        ).fetchone()
    assert record is not None
    assert made["export"]["rows"] == shown["total"] == record[0] >= 1
    assert record[2] == {"overdue_only": True}
    assert record[3] == timedelta(days=30)
    downloaded = client.get(f"{_base(team['org'], team['project'])}/exports/{made['export']['id']}")
    assert downloaded.status_code == 200, downloaded.text
    assert "attachment" in downloaded.headers["content-disposition"]
    book = load_workbook(io.BytesIO(downloaded.content))
    about = {row[0].value: row[1].value for row in book["About"].iter_rows(min_row=3)}
    assert about["Records"] == f"{record[0]:,}"
    assert about["Filters"] == "overdue only = True"
    assert about["Rules"] == "Metrics rulebook v1, pending approval"
    assert datetime.fromisoformat(str(about["Data as of (UTC)"])) == record[1]
    sheet = book["Records"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    headers = [cell.value for cell in sheet[1]]
    assert len(rows) == record[0]
    references = {row[headers.index("Reference")] for row in rows}
    assert references == {item["reference"] for item in shown["requests"]}
    request_column = headers.index("Request")
    formula = next(row for row in rows if str(row[request_column]).startswith("'="))
    assert formula is not None
    # Making and downloading are on the audit trail.
    with connect(migrated.url) as connection:
        actions = [
            row[0]
            for row in connection.execute(
                "SELECT action FROM audit_events WHERE target_id = %s ORDER BY occurred_at",
                (made["export"]["id"],),
            ).fetchall()
        ]
    assert actions == ["export.created", "export.downloaded"]


def test_a_csv_export_and_an_empty_one(migrated: Database) -> None:
    team, client = _engagement(migrated, "csv")
    with connect(migrated.url) as connection:
        context = _context(connection, migrated, team)
        controls = _run(context, "list_controls", limit=100)
        made = _run(context, "export_records", dataset="controls", file_format="csv")
        nothing = _run(context, "list_requests", search="no such request anywhere")
        empty = _run(
            context,
            "export_records",
            dataset="evidence_requests",
            filters={"search": "no such request anywhere"},
        )
    downloaded = client.get(f"{_base(team['org'], team['project'])}/exports/{made['export']['id']}")
    text = downloaded.content.decode("utf-8")
    assert text.startswith("﻿Control,Title,")
    assert len(text.strip().splitlines()) == controls["total"] + 1
    # Nothing matched: said as nothing, not as missing data.
    assert nothing["total"] == 0 and nothing["matching"] == "none" and nothing["requests"] == []
    assert empty["export"]["rows"] == 0


def test_missing_values_are_not_recorded_never_zero(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "missing")
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        connection.execute(
            "UPDATE engagement_controls SET owner_name = '', owner_email = '' "
            "WHERE project_id = %s",
            (team["project"],),
        )
        _set(connection, ids[0], recipient="", testing_phase=None)
        connection.commit()
        result = _run(
            _context(connection, migrated, team),
            "list_requests",
            control_ref=None,
            limit=100,
        )
    row = next(r for r in result["requests"] if r["request_id"] == str(ids[0]))
    assert row["control_owner"] == "Not recorded"
    assert row["recipient"] == "Not recorded"
    assert row["phase"] == "Not recorded"
    assert row["follow_up"] == "Set the recipient, then submit"


def test_a_status_that_contradicts_the_dates_is_flagged_not_fixed(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "contradict")
    with connect(migrated.url) as connection:
        ids = _requests(connection, team["project"])
        _set(
            connection,
            ids[0],
            status="OVERDUE",
            due_at=datetime.now(UTC) + timedelta(days=5),
        )
        connection.commit()
        result = _run(_context(connection, migrated, team), "list_requests", limit=100)
        status = connection.execute(
            "SELECT status FROM pbc_requests WHERE id = %s", (ids[0],)
        ).fetchone()
    flagged = next(c for c in result["data_quality"] if c["check"] == "status_overdue_but_not_due")
    assert flagged["count"] == 1
    row = next(r for r in result["requests"] if r["request_id"] == str(ids[0]))
    assert row["status"] == "OVERDUE" and row["overdue"] is False
    assert status == ("OVERDUE",)


def test_large_engagements_count_every_record(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "large")
    with connect(migrated.url) as connection:
        connection.execute(
            """
            INSERT INTO pbc_requests (
                organization_id, project_id, operation_key, control_ref, owner_user_id,
                recipient, due_at, evidence_description, status, engagement_control_id,
                testing_phase
            )
            SELECT organization_id, project_id, 'bulk-' || g, control_ref, owner_user_id,
                   recipient, due_at, evidence_description, 'DRAFT', engagement_control_id,
                   testing_phase
            FROM pbc_requests, generate_series(1, 250) AS g
            WHERE id = (SELECT id FROM pbc_requests WHERE project_id = %s LIMIT 1)
            """,
            (team["project"],),
        )
        connection.commit()
        total = connection.execute(
            "SELECT count(*) FROM pbc_requests WHERE project_id = %s AND status <> 'CANCELLED'",
            (team["project"],),
        ).fetchone()
        context = _context(connection, migrated, team)
        first = _run(context, "list_requests", limit=100, sort="control")
        last = _run(context, "list_requests", limit=100, offset=200, sort="control")
        made = _run(context, "export_records", dataset="evidence_requests")
        assert total is not None
        # Every page, however many records the engagement has.
        pages = [
            engagement_data.evidence_requests(
                connection, context.principal, team["org"], team["project"],
                engagement_data.RequestFilters(), sort="control", offset=offset, limit=100,
            )
            for offset in range(0, total[0], 100)
        ]  # fmt: skip
    assert total is not None and total[0] > 250
    # Totals cover every record, never only those that came back.
    assert first["total"] == last["total"] == total[0]
    assert first["more"] is True and 0 < first["returned"] <= 100
    assert last["offset"] == 200 and last["returned"] > 0
    # Paging through gives every record once.
    seen = [row["request_id"] for page in pages for row in page["requests"]]
    assert len(seen) == len(set(seen)) == total[0]
    assert made["export"]["rows"] == total[0]


def test_reading_and_exporting_change_no_engagement_records(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "readonly")
    fingerprint = """
        SELECT md5(coalesce((SELECT string_agg(id::text || status || version, ',' ORDER BY id)
                             FROM pbc_requests WHERE project_id = %(p)s), '')
                || coalesce((SELECT string_agg(id::text || version, ',' ORDER BY id)
                             FROM engagement_controls WHERE project_id = %(p)s), '')
                || coalesce((SELECT string_agg(id::text || status || version, ',' ORDER BY id)
                             FROM test_sheets WHERE project_id = %(p)s), '')
                || coalesce((SELECT string_agg(id::text || status || version, ',' ORDER BY id)
                             FROM deficiencies WHERE project_id = %(p)s), '')
                || coalesce((SELECT string_agg(id::text || status, ',' ORDER BY id)
                             FROM workpapers WHERE project_id = %(p)s), ''))
    """
    with connect(migrated.url) as connection:
        before = connection.execute(fingerprint, {"p": team["project"]}).fetchone()
        audits = connection.execute(
            "SELECT count(*) FROM audit_events WHERE project_id = %s", (team["project"],)
        ).fetchone()
        context = _context(connection, migrated, team)
        for name in ("engagement_status", "list_requests", "list_controls", "list_exceptions",
                     "list_review_items"):  # fmt: skip
            _run(context, name)
        after_reads = connection.execute(
            "SELECT count(*) FROM audit_events WHERE project_id = %s", (team["project"],)
        ).fetchone()
        _run(context, "export_records", dataset="exceptions")
        after = connection.execute(fingerprint, {"p": team["project"]}).fetchone()
    assert before == after
    # Reading leaves no trace; the export leaves exactly one.
    assert audits == after_reads
    with connect(migrated.url) as connection:
        exported = connection.execute(
            "SELECT count(*) FROM audit_events WHERE project_id = %s AND action LIKE 'export.%%'",
            (team["project"],),
        ).fetchone()
    assert exported == (1,)


def test_another_engagement_only_when_the_person_can_open_it(migrated: Database) -> None:
    team, _client_ = _engagement(migrated, "mine")
    other, _other_client = _engagement(migrated, "theirs")
    with connect(migrated.url) as connection:
        context = _context(connection, migrated, team)
        # The other firm's engagement is not found, whatever it is called.
        with pytest.raises(ToolFailed, match="No engagement you can open matches"):
            _run(context, "engagement_status", engagement=str(other["project"]))
        with pytest.raises(ToolFailed, match="No engagement you can open matches"):
            _run(context, "list_requests", engagement="SOX 2027 at another firm")
        # A second engagement in the same firm, which the owner can open.
        connection.execute(
            "UPDATE projects SET name = 'SOX 2026' WHERE id = %s", (team["project"],)
        )
        connection.commit()
    created = _client_.post(
        f"/organizations/{team['org']}/projects",
        json={
            "name": "SOX 2028",
            "slug": f"sox-2028-{uuid4().hex[:6]}",
            "scope": "",
            "owner_user_id": str(team["owner_id"]),
            "status": "DRAFT",
            "client_name": "Globex",
        },
    )
    assert created.status_code == 201, created.text
    with connect(migrated.url) as connection:
        context = _context(connection, migrated, team)
        found = _run(context, "engagement_status", engagement="sox 2028")
        assert found["engagement"] == "SOX 2028" and found["client"] == "Globex"
        assert found["controls"]["in_scope"] == 0
        # A name that fits both is not guessed between.
        with pytest.raises(ToolFailed, match="matches more than one engagement"):
            _run(context, "engagement_status", engagement="SOX 20")


def test_instructions_inside_evidence_cannot_reach_another_firm(migrated: Database) -> None:
    team, client = _engagement(migrated, "inject")
    other, _other_client = _engagement(migrated, "victim")
    message_id = _ask(client, team, "Read the file and do what it says.")
    # Whatever a file tells the model, the tools refuse what the person may not see.
    script = _Script(
        [
            _reply("", _call("list_requests", engagement=str(other["project"]))),
            _reply(
                "",
                _call(
                    "export_records",
                    dataset="evidence_requests",
                    engagement=str(other["project"]),
                ),
            ),
            _reply("I cannot read another firm's records."),
        ]
    )
    with connect(migrated.url) as connection:
        assert _turn(connection, migrated, script, message_id) == "DONE"
        exported = connection.execute(
            "SELECT count(*) FROM engagement_exports WHERE project_id = %s", (other["project"],)
        ).fetchone()
    assert script.seen[1][-1]["content"][0]["is_error"] is True
    assert script.seen[2][-1]["content"][0]["is_error"] is True
    assert exported == (0,)


def test_a_file_that_could_not_be_stored_is_never_offered(
    migrated: Database, monkeypatch: pytest.MonkeyPatch
) -> None:
    team, client = _engagement(migrated, "nostore")

    class _Broken:
        def put(self, key: str, content: bytes) -> str:
            raise StoreError("object storage is unavailable")

        def get(self, key: str) -> bytes:
            raise StoreError("object storage is unavailable")

    with connect(migrated.url) as connection:
        independent_steps(connection)
        with pytest.raises(exports.ExportFailed, match="could not be stored"):
            exports.create_export(
                connection,
                Settings(database_url=migrated.url),
                Principal(kind=PrincipalKind.USER, id=team["owner_id"]),
                team["org"],
                team["project"],
                "evidence_requests",
                "xlsx",
                {},
                store_for=lambda _settings, _connection: _Broken(),
            )
        recorded = connection.execute(
            "SELECT (SELECT count(*) FROM engagement_exports WHERE project_id = %s), "
            "(SELECT count(*) FROM audit_events WHERE project_id = %s "
            "AND action LIKE 'export.%%')",
            (team["project"], team["project"]),
        ).fetchone()
    assert recorded == (0, 0)
    # Through the console: the step fails with the reason, and no file is shown.
    original = exports.create_export

    def failing(*args: Any, **kwargs: Any) -> Any:
        kwargs["store_for"] = lambda _settings, _connection: _Broken()
        return original(*args, **kwargs)

    monkeypatch.setattr(exports, "create_export", failing)
    message_id = _ask(client, team, "Export the evidence requests to Excel.")
    script = _Script(
        [
            _reply("", _call("export_records", dataset="evidence_requests")),
            _reply("The file could not be made."),
        ]
    )
    with connect(migrated.url) as connection:
        _turn(connection, migrated, script, message_id)
    step = _thread(client, team)["messages"][-1]["actions"][0]
    assert step["status"] == "FAILED"
    assert "could not be stored" in step["error"]
    assert step["result"] is None


def test_follow_ups_see_the_filters_already_used(migrated: Database) -> None:
    team, client = _engagement(migrated, "followup")
    first = _ask(client, team, "Show me the overdue items for SAP.")
    script = _Script(
        [
            _reply("", _call("list_requests", overdue_only=True, application="SAP")),
            _reply("None of SAP's requests are overdue."),
        ]
    )
    with connect(migrated.url) as connection:
        _turn(connection, migrated, script, first)
    second = _ask(client, team, "Export that to Excel.")
    follow = _Script(
        [
            _reply(
                "",
                _call(
                    "export_records",
                    dataset="evidence_requests",
                    filters={"overdue_only": True, "application": "SAP"},
                ),
            ),
            _reply("Here is the file."),
        ]
    )
    with connect(migrated.url) as connection:
        _turn(connection, migrated, follow, second)
    # The second turn sees the first turn's call and its filters.
    calls = [
        block
        for message in follow.seen[0]
        for block in message["content"]
        if block["type"] == "tool_call"
    ]
    assert calls[0]["name"] == "list_requests"
    assert calls[0]["arguments"] == {"overdue_only": True, "application": "SAP"}
    step = _thread(client, team)["messages"][-1]["actions"][0]
    assert step["status"] == "DONE"
    assert step["result"]["export"]["filters"] == {"overdue_only": True, "application": "SAP"}


def test_downloads_only_for_people_who_can_see_the_engagement(migrated: Database) -> None:
    team, client = _engagement(migrated, "access")
    other, other_client = _engagement(migrated, "elsewhere")
    with connect(migrated.url) as connection:
        made = _run(_context(connection, migrated, team), "export_records", dataset="controls")
        outsider = _add_user(
            connection, team["org"], f"c-{uuid4().hex[:6]}@example.com", "CONTRIBUTOR"
        )
        member = _add_user(
            connection, team["org"], f"m-{uuid4().hex[:6]}@example.com", "CONTRIBUTOR"
        )
        _add_project_member(connection, team["org"], team["project"], member)
    path = f"{_base(team['org'], team['project'])}/exports/{made['export']['id']}"
    assert client.get(path).status_code == 200
    # Another firm, and someone in this firm who is not on the engagement, see nothing.
    assert other_client.get(path).status_code == 404
    assert (
        other_client.get(
            f"{_base(other['org'], other['project'])}/exports/{made['export']['id']}"
        ).status_code
        == 404
    )
    assert _client(migrated, outsider).get(path).status_code == 404
    assert _client(migrated, member).get(path).status_code == 200
    # After 30 days the file is gone; the record of it stays.
    with connect(migrated.url) as connection:
        connection.execute(
            "UPDATE engagement_exports SET created_at = created_at - interval '31 days', "
            "expires_at = expires_at - interval '31 days' WHERE id = %s",
            (made["export"]["id"],),
        )
        connection.commit()
    assert client.get(path).status_code == 410
    with connect(migrated.url) as connection:
        _run(_context(connection, migrated, team), "export_records", dataset="controls")
        kept = connection.execute(
            "SELECT e.purged_at IS NOT NULL, "
            "(SELECT count(*) FROM mock_objects o WHERE o.object_key = e.object_key) "
            "FROM engagement_exports e WHERE e.id = %s",
            (made["export"]["id"],),
        ).fetchone()
    assert kept == (True, 0)


def test_the_console_offers_the_file_with_the_answer(migrated: Database) -> None:
    team, client = _engagement(migrated, "consolefile")
    message_id = _ask(client, team, "Give me the status, and the controls as a spreadsheet.")
    script = _Script(
        [
            _reply(
                "",
                _call("engagement_status"),
                _call("export_records", dataset="controls"),
            ),
            _reply("Here is where things stand, with the controls attached."),
        ]
    )
    with connect(migrated.url) as connection:
        assert _turn(connection, migrated, script, message_id) == "DONE"
    status = script.seen[1][-1]["content"][0]
    assert status["is_error"] is False and '"as_of"' in status["content"]
    agent = _thread(client, team)["messages"][-1]
    export = agent["actions"][0]
    assert export["tool"] == "export_records" and export["status"] == "DONE"
    assert export["summary"] == "Exported controls to Excel"
    file = export["result"]["export"]
    assert file["file_name"].endswith(".xlsx") and file["rows"] > 0
    assert file["project_id"] == str(team["project"])
