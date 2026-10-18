"""Workpapers: filled from the records, drafted by the agent, edited, versioned, exported."""

from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from support import Database
from test_evidence_analysis import _Provider
from test_evidence_intake import _accept_population, _engagement, _request
from test_outbound import migrated
from test_testing_engine import _add, _answer, _base, _Citing, _read_all, _sheet, _year

from ai_workforce_api.db import connect
from ai_workforce_api.testing_engine import run_sheet
from ai_workforce_api.workpaper_service import run_draft

_ = migrated


def _tested_control(
    migrated: Database, slug: str, sign: bool = True
) -> tuple[dict[str, Any], Any, Any, str]:
    """A sampled control, tested with one exception and (unless told not to) signed off."""
    team, client, scheduler = _engagement(migrated, slug)
    base = _base(team)
    kickoff = _request(migrated, team, "ITGC-AC-01")
    year = _year(migrated, team)
    population = "User ID,Name,Date granted\n" + "".join(
        f"u000{n},Demo User {n},{year}-0{n}-10\n" for n in range(1, 4)
    )
    _add(client, team, "access population.csv", population, kickoff["id"])
    _read_all(migrated, scheduler)
    plan = client.get(f"{base}/control-tests").json()[0]
    _accept_population(client, team, plan["engagement_control_id"])
    approved = client.post(
        f"{base}/control-tests/{plan['id']}/approve", json={"version": plan["version"]}
    ).json()
    sheet = _sheet(client, team)
    labels = [item["label"] for item in sheet["items"]]
    tickets = "ticket,user,approved_by\n" + "".join(
        f"T-{n},{label},Demo Manager\n" for n, label in enumerate(labels)
    )
    _add(client, team, "access tickets.csv", tickets, approved["sample_request_id"])
    _read_all(migrated, scheduler)
    attributes = len(sheet["attributes"])
    provider = _Citing(
        [
            _answer(["PASS"] * attributes),
            _answer(["EXCEPTION"] + ["PASS"] * (attributes - 1)),
            _answer(["PASS"] * attributes),
        ]
    )
    with connect(migrated.url) as connection:
        run_sheet(connection, provider, UUID(sheet["id"]))
    sheet = _sheet(client, team)
    if sign:
        signed = client.post(
            f"{base}/test-sheets/{sheet['id']}/sign-off", json={"version": sheet["version"]}
        )
        assert signed.status_code == 200, signed.text
    files = client.get(f"{base}/evidence-register").json()["files"]
    ticket_file = next(f["id"] for f in files if f["file_name"] == "access tickets.csv")
    # Later steps read more files with the same scheduler.
    team["scheduler"] = scheduler
    return team, client, sheet, ticket_file


def test_a_signed_off_test_becomes_a_drafted_effectiveness_workpaper(migrated: Database) -> None:
    team, client, sheet, ticket_file = _tested_control(migrated, "wp")
    base = f"{_base(team)}/workpapers"
    listed = client.get(base, params={"control_id": sheet["control_id"]}).json()
    assert [(w["stage"], w["drafting"], w["current_version"]) for w in listed] == [
        ("EFFECTIVENESS", True, 0)
    ]
    workpaper = listed[0]
    draft = {
        "sections": [
            {
                "key": "population",
                "text": "The population came from the user listing. Tickets were provided "
                f'[cite:{ticket_file} "ticket | user | approved_by"].',
            },
            # The conclusion is derived from the testing, never the agent's.
            {"key": "interim-conclusion", "text": "The control is effective."},
        ]
    }
    provider = _Provider([draft])
    with connect(migrated.url) as connection:
        assert run_draft(connection, provider, UUID(workpaper["id"])) == "DONE"
    assert "Draft these workpaper sections." in provider.requests[0].prompt
    body = client.get(base, params={"control_id": sheet["control_id"]}).json()[0]
    assert (body["current_version"], body["drafting"]) == (1, False)
    content = body["content"]
    assert content["reference"] == "ITGC-AC-01-SE-TOE"
    sections = {s["key"]: s for s in content["sections"]}
    population = sections["population"]
    assert population["source"] == "AGENT"
    assert population["text"].endswith("Tickets were provided [^1].")
    assert population["citations"][0]["quote"] == "ticket | user | approved_by"
    assert sections["interim-conclusion"]["source"] == "RECORDS"
    assert "did not operate effectively" in sections["interim-conclusion"]["text"]
    # One workpaper holds the interim and roll-forward tabs; the roll-forward's own
    # sections wait for its round.
    assert [t["key"] for t in content["tabs"]] == ["interim", "rollforward"]
    assert {s["tab"] for s in content["sections"]} == {"interim", "rollforward"}
    assert sections["rollforward-exceptions"]["tab"] == "rollforward"
    assert "No roll-forward round has started" in sections["changes-since-interim"]["text"]
    fields = {f["label"]: f["value"] for f in sections["sample-selection"]["fields"]}
    assert fields["Sample size"] and fields["Selection method"].startswith("Random, seed")
    matrix = sections["testing-matrix"]
    assert len(matrix["rows"]) == 3
    assert sum(1 for row in matrix["rows"] if any(c.startswith("Exception") for c in row)) == 1
    assert len(sections["exceptions"]["rows"]) == 1
    assert sections["sign-off"]["rows"][0][0] == "Testing signed off by"

    # A person edits the conclusion; an edit on an old version is refused.
    path = f"{base}/{workpaper['id']}"
    edited = client.post(
        f"{path}/sections/interim-conclusion",
        json={"based_on": 1, "text": '=HYPERLINK("http://x") Exceptions in one sample.'},
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["current_version"] == 2
    stale = client.post(
        f"{path}/sections/interim-conclusion", json={"based_on": 1, "text": "Again"}
    )
    assert stale.status_code == 409
    # Drafting again keeps the person's section.
    again = client.post(f"{base}/draft", json={"control_id": sheet["control_id"]})
    assert again.status_code == 200
    with connect(migrated.url) as connection:
        run_draft(connection, _Provider([draft]), UUID(workpaper["id"]))
    latest = client.get(base, params={"control_id": sheet["control_id"]}).json()
    effectiveness = next(w for w in latest if w["stage"] == "EFFECTIVENESS")
    conclusion = next(
        s for s in effectiveness["content"]["sections"] if s["key"] == "interim-conclusion"
    )
    assert conclusion["source"] == "PERSON"
    assert effectiveness["current_version"] == 3
    assert [v["source"] for v in effectiveness["versions"]] == ["AGENT", "PERSON", "AGENT"]
    first = client.get(f"{path}/versions/1").json()
    assert first["shown_version"] == 1
    # The design workpaper opened too, and is drafted on request.
    design = next(w for w in latest if w["stage"] == "DESIGN")
    with connect(migrated.url) as connection:
        run_draft(connection, None, UUID(design["id"]))
    # Prepared, the workpaper is frozen until reopened.
    current = client.get(base, params={"control_id": sheet["control_id"]}).json()
    effectiveness = next(w for w in current if w["stage"] == "EFFECTIVENESS")
    prepared = client.post(f"{path}/prepare", json={"version": effectiveness["version"]})
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()["status"] == "PREPARED"
    frozen = client.post(f"{path}/sections/conclusion", json={"based_on": 3, "text": "x"})
    assert frozen.status_code == 409

    # One workpaper, and the binder with an index that links to each sheet.
    one = client.get(f"{path}/export.xlsx")
    assert one.status_code == 200
    assert one.headers["content-disposition"].endswith('ITGC-AC-01-SE-TOE.xlsx"')
    book = load_workbook(BytesIO(one.content))
    values = [str(c.value) for row in book.active.iter_rows() for c in row if c.value]
    assert "Testing matrix" in values
    assert any(v.startswith("Prepared by") for v in values)
    # A formula typed into a section is written as text, never as a formula.
    assert '\'=HYPERLINK("http://x") Exceptions in one sample.' in values
    assert any(v.startswith("[1] access tickets.csv") for v in values)
    binder = client.get(f"{base}/export.xlsx")
    assert binder.status_code == 200
    workbook = load_workbook(BytesIO(binder.content))
    # One sheet per tab: design, then the interim and roll-forward tabs of the testing.
    assert workbook.sheetnames == [
        "Index",
        "ITGC-AC-01-SE-TOD",
        "ITGC-AC-01-SE-TOE",
        "ITGC-AC-01-SE-RF",
    ]
    index = workbook["Index"]
    assert index["A4"].value == "ITGC-AC-01-SE-TOD"
    assert index["C6"].value == "IT general control: Roll-forward"
    assert index["F5"].hyperlink is not None
    rollforward = [
        str(c.value) for row in workbook["ITGC-AC-01-SE-RF"].iter_rows() for c in row if c.value
    ]
    assert "Changes since interim" in rollforward and "Testing matrix" not in rollforward
    with connect(migrated.url) as connection:
        audit = str(
            connection.execute(
                "SELECT array_agg(action) FROM audit_events WHERE target_type = 'workpaper'"
            ).fetchone()
        )
    for action in ("workpaper.drafted", "workpaper.edited", "workpaper.prepared"):
        assert action in audit


def test_an_effectiveness_workpaper_needs_signed_off_testing(migrated: Database) -> None:
    team, client, _scheduler = _engagement(migrated, "wpearly")
    base = f"{_base(team)}/workpapers"
    controls = client.get(f"{_base(team)}/workspace/controls").json()
    control = next(c for c in controls if c["control_ref"] == "ITGC-AC-02")
    opened = client.post(f"{base}/draft", json={"control_id": control["id"]})
    assert opened.status_code == 200, opened.text
    papers = {w["stage"]: w for w in opened.json()}
    with connect(migrated.url) as connection:
        for paper in papers.values():
            run_draft(connection, None, UUID(paper["id"]))
    papers = {w["stage"]: w for w in client.get(base, params={"control_id": control["id"]}).json()}
    early = client.post(
        f"{base}/{papers['EFFECTIVENESS']['id']}/prepare",
        json={"version": papers["EFFECTIVENESS"]["version"]},
    )
    assert early.status_code == 409
    assert "testing" in early.json()["detail"]
    assert "No model is configured" in papers["DESIGN"]["draft_note"]
    design = client.post(
        f"{base}/{papers['DESIGN']['id']}/prepare", json={"version": papers["DESIGN"]["version"]}
    )
    assert design.status_code == 200


def test_the_testers_first_results_start_the_effectiveness_workpaper(
    migrated: Database,
) -> None:
    """The workpaper is built while testing is open, from results marked as proposed;
    it cannot be prepared until a person signs the testing off, and the sign-off asks for
    a redraft from the confirmed results."""
    from ai_workforce_api.workpaper_service import draft_effectiveness_after_run

    team, client, sheet, _ticket = _tested_control(migrated, "wpprog", sign=False)
    base = f"{_base(team)}/workpapers"
    assert sheet["status"] == "OPEN" and sheet["run_state"] == "IDLE"
    assert client.get(base, params={"control_id": sheet["control_id"]}).json() == []
    with connect(migrated.url) as connection:
        queued = draft_effectiveness_after_run(connection, UUID(sheet["id"]))
        assert len(queued) == 1
        # Later runs do not queue it again.
        assert draft_effectiveness_after_run(connection, UUID(sheet["id"])) == []
    listed = client.get(base, params={"control_id": sheet["control_id"]}).json()
    assert [(w["stage"], w["drafting"], w["draft_state"]) for w in listed] == [
        ("EFFECTIVENESS", True, "QUEUED")
    ]
    provider = _Provider([{"sections": [{"key": "population", "text": "Users were listed."}]}])
    with connect(migrated.url) as connection:
        assert run_draft(connection, provider, UUID(listed[0]["id"])) == "DONE"
    request = provider.requests[0]
    assert '"status": "OPEN"' in request.prompt
    assert "proposals that no person has confirmed" in request.system
    paper = client.get(base, params={"control_id": sheet["control_id"]}).json()[0]
    assert (paper["draft_state"], paper["current_version"]) == ("IDLE", 1)
    sections = {s["key"]: s for s in paper["content"]["sections"]}
    assert sections["population"]["source"] == "AGENT"
    # The conclusion is not the agent's, and not drawn before the sign-off.
    assert sections["interim-conclusion"]["source"] == "RECORDS"
    assert "did not operate effectively" not in sections["interim-conclusion"]["text"]
    # The gate holds: no preparing before the testing is signed off.
    early = client.post(f"{base}/{paper['id']}/prepare", json={"version": paper["version"]})
    assert early.status_code == 409 and "testing" in early.json()["detail"]
    # The sign-off asks for a redraft from the confirmed results.
    sheet = _sheet(client, team)
    signed = client.post(
        f"{_base(team)}/test-sheets/{sheet['id']}/sign-off", json={"version": sheet["version"]}
    )
    assert signed.status_code == 200, signed.text
    paper = client.get(base, params={"control_id": sheet["control_id"]}).json()[0]
    assert (paper["drafting"], paper["draft_state"]) == (True, "QUEUED")


def test_no_effectiveness_draft_before_the_tester_has_results(migrated: Database) -> None:
    from ai_workforce_api.workpaper_service import draft_effectiveness_after_run

    team, client, _scheduler = _engagement(migrated, "wpnoresult")
    controls = client.get(f"{_base(team)}/workspace/controls").json()
    control = next(c for c in controls if c["control_ref"] == "ITGC-AC-02")
    with connect(migrated.url) as connection:
        sheet = connection.execute(
            """
            INSERT INTO test_sheets (organization_id, project_id, engagement_control_id,
                                     method, attributes, requested_by)
            VALUES (%s, %s, %s, 'TEST_OF_ONE', '["Approved"]', %s)
            RETURNING id
            """,
            (team["org"], team["project"], control["id"], team["owner_id"]),
        ).fetchone()
        assert sheet is not None
        connection.execute(
            "INSERT INTO test_items (test_sheet_id, position, label) VALUES (%s, 1, 'One')",
            (sheet[0],),
        )
        connection.commit()
        assert draft_effectiveness_after_run(connection, UUID(str(sheet[0]))) == []


def test_walkthrough_cells_write_out_checked_citations(migrated: Database) -> None:
    team, client, sheet, ticket_file = _tested_control(migrated, "wpwalk")
    base = f"{_base(team)}/workpapers"
    assert client.post(f"{base}/draft", json={"control_id": sheet["control_id"]}).status_code == 200
    listed = client.get(base, params={"control_id": sheet["control_id"]}).json()
    design = next(w for w in listed if w["stage"] == "DESIGN")
    made_up = "00000000-0000-7000-8000-000000000000"
    draft = {
        "sections": [],
        "walkthrough": [
            {
                "step": "Access requested and approved",
                "who": "Demo Manager",
                "evidence": f'[cite:{ticket_file} "ticket | user | approved_by"]'
                f' [cite:{made_up} "words nobody wrote"]',
                "result": "Approved before access",
            }
        ],
    }
    with connect(migrated.url) as connection:
        assert run_draft(connection, _Provider([draft]), UUID(design["id"])) == "DONE"
    body = client.get(base, params={"control_id": sheet["control_id"]}).json()
    paper = next(w for w in body if w["stage"] == "DESIGN")
    walkthrough = next(
        s for s in paper["content"]["sections"] if s["key"] == "walkthrough-of-one-instance"
    )
    assert walkthrough["source"] == "AGENT"
    evidence = walkthrough["rows"][0][2]
    # A table has no footnotes: the checked passage is written out, the made-up one dropped.
    assert evidence.startswith("access tickets.csv")
    assert "“ticket | user | approved_by”" in evidence
    assert "[cite:" not in evidence and "nobody" not in evidence


def test_the_agent_waits_for_evidence_before_drafting(migrated: Database) -> None:
    """A model is configured, but the control has no read evidence and no signed-off
    testing: the records are filled, the model is not called, and the draft says why."""
    team, client, _scheduler = _engagement(migrated, "wpwait")
    base = f"{_base(team)}/workpapers"
    controls = client.get(f"{_base(team)}/workspace/controls").json()
    control = next(c for c in controls if c["control_ref"] == "ITGC-AC-02")
    opened = client.post(f"{base}/draft", json={"control_id": control["id"]})
    assert opened.status_code == 200, opened.text
    provider = _Provider(
        [{"sections": [{"key": "understanding-of-the-control", "text": "Made up."}]}]
    )
    with connect(migrated.url) as connection:
        for paper in opened.json():
            assert run_draft(connection, provider, UUID(paper["id"])) == "DONE"
    assert provider.requests == []
    papers = {w["stage"]: w for w in client.get(base, params={"control_id": control["id"]}).json()}
    for paper in papers.values():
        assert paper["drafting"] is False
        assert "Waiting for evidence" in paper["draft_note"]
        assert all(s["source"] != "AGENT" for s in paper["content"]["sections"])


def test_the_first_read_file_opens_the_design_workpaper(migrated: Database) -> None:
    from ai_workforce_api.workpaper_service import draft_design_after_analysis

    team, client, scheduler = _engagement(migrated, "wpfirst")
    base = f"{_base(team)}/workpapers"
    kickoff = _request(migrated, team, "ITGC-AC-01")
    controls = client.get(f"{_base(team)}/workspace/controls").json()
    control_id = next(c["id"] for c in controls if c["control_ref"] == "ITGC-AC-01")
    _add(client, team, "access policy.txt", "New users are approved by the manager.", kickoff["id"])
    assert client.get(base, params={"control_id": control_id}).json() == []
    with connect(migrated.url) as connection:
        # Not read yet: nothing to draft from.
        assert draft_design_after_analysis(connection, scheduler.analyses[0]) == []
        _read_all(migrated, scheduler)
        queued = draft_design_after_analysis(connection, scheduler.analyses[0])
        assert len(queued) == 1
        # The next file, or a second read, does not queue it again.
        assert draft_design_after_analysis(connection, scheduler.analyses[0]) == []
    listed = client.get(base, params={"control_id": control_id}).json()
    assert [(w["stage"], w["drafting"]) for w in listed] == [("DESIGN", True)]
    provider = _Provider(
        [
            {
                "sections": [
                    {"key": "understanding-of-the-control", "text": "Managers approve new users."}
                ]
            }
        ]
    )
    with connect(migrated.url) as connection:
        assert run_draft(connection, provider, queued[0]) == "DONE"
    assert "access policy.txt" in provider.requests[0].prompt
    design = client.get(base, params={"control_id": control_id}).json()[0]
    sections = {s["key"]: s for s in design["content"]["sections"]}
    assert sections["understanding-of-the-control"]["source"] == "AGENT"
    assert design["draft_note"] == ""


def test_the_local_worker_starts_the_testing_workpaper_after_a_run(migrated: Database) -> None:
    """The worker's own chain: a finished run queues the effectiveness draft, once."""
    from ai_workforce_api.pbc_drafting import LocalDraftScheduler
    from ai_workforce_api.settings import Settings

    team, client, sheet, _ticket = _tested_control(migrated, "wpchain", sign=False)
    scheduled: list[list[UUID]] = []
    worker = LocalDraftScheduler(Settings(database_url=migrated.url), test=lambda _id: "DONE")
    worker.schedule_workpaper_drafts = scheduled.append  # type: ignore[method-assign,assignment]
    try:
        worker._attempt_test(UUID(sheet["id"]))
        worker._attempt_test(UUID(sheet["id"]))
    finally:
        worker.shutdown()
    listed = client.get(
        f"{_base(team)}/workpapers", params={"control_id": sheet["control_id"]}
    ).json()
    assert [w["stage"] for w in listed] == ["EFFECTIVENESS"]
    # Queued after the first run; the second run found it open already.
    assert scheduled == [[UUID(listed[0]["id"])], []]
